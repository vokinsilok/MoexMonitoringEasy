from __future__ import annotations

import asyncio
from dataclasses import asdict
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import html
import json
import re
import tempfile
from pathlib import Path

from aiogram import Bot, F, Router
from aiogram.exceptions import TelegramBadRequest, TelegramNetworkError
from aiogram.filters import Command, CommandStart, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.types import (
    CallbackQuery,
    FSInputFile,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
)

from bot_service.keyboards import (
    CB_BACK_TO_LIST,
    CB_FAVORITE_TOGGLE_PREFIX,
    CB_MODE_ALL,
    CB_MODE_FAVORITES,
    CB_PREFIX_DETAILS,
    CB_PREFIX_PAGE,
    CB_REFRESH,
    CB_REFRESH_DETAILS,
    CB_SEARCH,
    CB_TYPE_ALL,
    CB_TYPE_BOND,
    CB_TYPE_ETF,
    CB_TYPE_SHARE,
    PAGE_SIZE,
    share_details_keyboard,
    shares_list_keyboard,
)
from bot_service.access import ADMIN_TELEGRAM_IDS, is_admin_user
from bot_service.models import ShareViewItem
from bot_service.service import TelegramSharesBrowserService
from bot_service.states import SharesBrowserStates


router = Router()
service = TelegramSharesBrowserService()
BACKGROUND_TASKS: set[asyncio.Task] = set()

# User-facing labels (kept explicit to avoid mojibake regressions).
MENU_SHARES = "📊 Инструменты"
MENU_NEW_MONITOR = "🔔 Новый мониторинг"
MENU_MONITORS = "🧭 Мои мониторинги"
MENU_CALENDAR = "🗓 Календарь"
MENU_PORTFOLIO = "💼 Портфель"
MENU_PORTFOLIO_ANALYSIS = "🤖 Анализ портфеля"
MENU_OPERATIONS = "🕘 Операции"
MENU_PROFILE = "🔐 Профиль T-Bank"
MENU_ORDER = "🧾 Заявка"
MENU_STOP = "🛑 Стоп-приказ"
MENU_ACTIVE_ORDERS = "📂 Активные заявки"
MENU_ADMIN = "🛡 Админка"

ORDER_TYPE_LIMIT_LABEL = "Лимитная"
ORDER_TYPE_MARKET_LABEL = "Рыночная"
ORDER_TYPE_BESTPRICE_LABEL = "Лучшая цена"

STOP_KIND_STOP_LOSS_LABEL = "Стоп-лосс"
STOP_KIND_STOP_LOSS_LIMIT_LABEL = "Стоп-лосс лимит"
STOP_KIND_TAKE_PROFIT_LABEL = "Тейк-профит"

DIRECTION_BUY_LABEL = "Покупка"
DIRECTION_SELL_LABEL = "Продажа"
CANCEL_TEXT = "Отмена"
MONITOR_INTERVAL_MIN_SECONDS = 5
MONITOR_INTERVAL_MAX_SECONDS = 86_400
MONITOR_INTERVAL_PRESETS_SECONDS: tuple[int, ...] = (15, 30, 60, 120, 300, 600)

MAIN_MENU_KEYBOARD = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text=MENU_SHARES), KeyboardButton(text=MENU_NEW_MONITOR)],
        [KeyboardButton(text=MENU_MONITORS), KeyboardButton(text=MENU_CALENDAR)],
        [KeyboardButton(text=MENU_PORTFOLIO), KeyboardButton(text=MENU_PORTFOLIO_ANALYSIS)],
        [KeyboardButton(text=MENU_OPERATIONS)],
        [KeyboardButton(text=MENU_PROFILE)],
        [KeyboardButton(text=MENU_ORDER), KeyboardButton(text=MENU_STOP)],
        [KeyboardButton(text=MENU_ACTIVE_ORDERS)],
    ],
    resize_keyboard=True,
)

MAIN_MENU_KEYBOARD_ADMIN = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text=MENU_SHARES), KeyboardButton(text=MENU_NEW_MONITOR)],
        [KeyboardButton(text=MENU_MONITORS), KeyboardButton(text=MENU_CALENDAR)],
        [KeyboardButton(text=MENU_PORTFOLIO), KeyboardButton(text=MENU_PORTFOLIO_ANALYSIS)],
        [KeyboardButton(text=MENU_OPERATIONS)],
        [KeyboardButton(text=MENU_PROFILE)],
        [KeyboardButton(text=MENU_ORDER), KeyboardButton(text=MENU_STOP)],
        [KeyboardButton(text=MENU_ACTIVE_ORDERS), KeyboardButton(text=MENU_ADMIN)],
    ],
    resize_keyboard=True,
)

MON_CB_LIST = "mon:list"
MON_CB_REFRESH = "mon:refresh"
MON_CB_ITEM_PREFIX = "mon:item"
MON_CB_ON_PREFIX = "mon:on"
MON_CB_OFF_PREFIX = "mon:off"
MON_CB_DEL_PREFIX = "mon:del"
MON_CB_REBASE_PREFIX = "mon:rebase"
MON_CB_EDIT_PREFIX = "mon:edit"
MON_CB_EDIT_PERCENT_PREFIX = "mon:edit_pct"
MON_CB_EDIT_RUB_PREFIX = "mon:edit_rub"
MON_CB_EDIT_BACK_PREFIX = "mon:edit_back"
CAL_CB_REFRESH = "cal:refresh"
CAL_CB_SETTINGS = "cal:settings"
CAL_CB_TOGGLE = "cal:toggle"
CAL_CB_DAYS_PREFIX = "cal:days"
PROF_CB_UPDATE = "prof:update"
PROF_CB_CHECK = "prof:check"
AO_CB_REFRESH = "ao:refresh"
AO_CB_CANCEL_ORDER_PREFIX = "ao:cancel_order"
AO_CB_CANCEL_STOP_PREFIX = "ao:cancel_stop"
ADM_CB_PANEL = "adm:panel"
ADM_CB_PENDING = "adm:pending"
ADM_CB_USERS = "adm:users"
ADM_CB_REFRESH = "adm:refresh"
ADM_CB_USER_PREFIX = "adm:user"
ADM_CB_APPROVE_PREFIX = "adm:approve"
ADM_CB_REVOKE_PREFIX = "adm:revoke"
PA_CB_PREFIX = "pa:h"


async def _safe_telegram_call(coro):
    try:
        return await coro
    except TelegramNetworkError:
        return None


async def _safe_edit_text(message: Message, text: str, reply_markup=None) -> None:
    try:
        await _safe_telegram_call(message.edit_text(text, reply_markup=reply_markup))
    except TelegramBadRequest as exc:
        if "message is not modified" not in str(exc).lower():
            raise


async def _safe_delete_message(message: Message | None) -> None:
    if message is None:
        return
    try:
        await message.delete()
    except Exception:
        return


async def _safe_delete_by_id(message: Message, message_id: int | None) -> None:
    if message_id is None:
        return
    try:
        await message.bot.delete_message(chat_id=message.chat.id, message_id=message_id)
    except Exception:
        return


async def _send_long_text(message: Message, text: str, *, chunk_size: int = 3600) -> None:
    if len(text) <= chunk_size:
        await _safe_telegram_call(message.answer(text))
        return
    start = 0
    while start < len(text):
        end = min(len(text), start + chunk_size)
        if end < len(text):
            split_at = text.rfind("\n", start, end)
            if split_at > start + 500:
                end = split_at
        await _safe_telegram_call(message.answer(text[start:end].strip()))
        start = end


def _menu_keyboard_for_user(telegram_user_id: int | None) -> ReplyKeyboardMarkup:
    if is_admin_user(telegram_user_id):
        return MAIN_MENU_KEYBOARD_ADMIN
    return MAIN_MENU_KEYBOARD


def _access_status_badge(status: str | None) -> str:
    normalized = str(status or "").strip().lower()
    if normalized == "approved":
        return "✅"
    if normalized == "pending":
        return "⏳"
    if normalized == "revoked":
        return "⛔"
    return "⚪"


def _display_user_name(item: dict) -> str:
    first_name = str(item.get("first_name") or "").strip()
    last_name = str(item.get("last_name") or "").strip()
    username = str(item.get("username") or "").strip()
    full_name = " ".join(part for part in [first_name, last_name] if part).strip()
    if full_name:
        return full_name
    if username:
        return f"@{username}"
    return "Без имени"


def _admin_panel_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="📥 Заявки", callback_data=ADM_CB_PENDING)],
            [InlineKeyboardButton(text="👥 Пользователи", callback_data=ADM_CB_USERS)],
            [InlineKeyboardButton(text="🔄 Обновить", callback_data=ADM_CB_REFRESH)],
        ]
    )


def _admin_user_actions_keyboard(telegram_user_id: int, status: str | None) -> InlineKeyboardMarkup:
    normalized = str(status or "").strip().lower()
    approve_text = "✅ Одобрить доступ" if normalized != "approved" else "✅ Доступ выдан"
    revoke_text = "⛔ Отозвать доступ" if normalized == "approved" else "⛔ Отклонить заявку"
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text=approve_text, callback_data=f"{ADM_CB_APPROVE_PREFIX}:{telegram_user_id}")],
            [InlineKeyboardButton(text=revoke_text, callback_data=f"{ADM_CB_REVOKE_PREFIX}:{telegram_user_id}")],
            [InlineKeyboardButton(text="📥 Заявки", callback_data=ADM_CB_PENDING)],
            [InlineKeyboardButton(text="👥 Пользователи", callback_data=ADM_CB_USERS)],
        ]
    )


def _admin_user_detail_text(item: dict) -> str:
    telegram_user_id = int(item.get("telegram_user_id") or 0)
    status = str(item.get("status") or "none")
    name = _display_user_name(item)
    username = str(item.get("username") or "").strip()
    username_text = f"@{username}" if username else "—"
    return (
        "🛡 <b>Пользователь бота</b>\n"
        f"ID: <code>{telegram_user_id}</code>\n"
        f"Имя: <b>{html.escape(name)}</b>\n"
        f"Username: <b>{html.escape(username_text)}</b>\n"
        f"Статус: <b>{_access_status_badge(status)} {html.escape(status)}</b>\n"
        f"Заявка: <b>{_format_datetime(str(item.get('requested_at') or ''))}</b>\n"
        f"Одобрен: <b>{_format_datetime(str(item.get('approved_at') or ''))}</b>\n"
        f"Отозван: <b>{_format_datetime(str(item.get('revoked_at') or ''))}</b>"
    )


def _admin_users_list_text(items: list[dict], *, title: str, total: int) -> str:
    lines = [f"🛡 <b>{title}</b>", f"Всего: <b>{total}</b>", ""]
    if not items:
        lines.append("Список пуст.")
        return "\n".join(lines)
    for idx, item in enumerate(items[:30], start=1):
        telegram_user_id = int(item.get("telegram_user_id") or 0)
        status = str(item.get("status") or "none")
        lines.append(
            f"{idx}. {_access_status_badge(status)} <code>{telegram_user_id}</code> · {html.escape(_display_user_name(item))}"
        )
    if len(items) > 30:
        lines.append("")
        lines.append("Показаны первые 30 записей.")
    lines.append("")
    lines.append("Нажмите на пользователя в кнопках ниже.")
    return "\n".join(lines)


def _admin_users_list_keyboard(items: list[dict]) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for item in items[:20]:
        telegram_user_id = int(item.get("telegram_user_id") or 0)
        if telegram_user_id <= 0:
            continue
        status = str(item.get("status") or "none")
        rows.append(
            [
                InlineKeyboardButton(
                    text=f"{_access_status_badge(status)} {telegram_user_id} · {_display_user_name(item)[:24]}",
                    callback_data=f"{ADM_CB_USER_PREFIX}:{telegram_user_id}",
                )
            ]
        )
    rows.append([InlineKeyboardButton(text="📥 Заявки", callback_data=ADM_CB_PENDING)])
    rows.append([InlineKeyboardButton(text="👥 Пользователи", callback_data=ADM_CB_USERS)])
    rows.append([InlineKeyboardButton(text="🔄 Обновить", callback_data=ADM_CB_REFRESH)])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _admin_request_keyboard(telegram_user_id: int) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="✅ Одобрить",
                    callback_data=f"{ADM_CB_APPROVE_PREFIX}:{telegram_user_id}",
                ),
                InlineKeyboardButton(
                    text="⛔ Отклонить",
                    callback_data=f"{ADM_CB_REVOKE_PREFIX}:{telegram_user_id}",
                ),
            ],
            [InlineKeyboardButton(text="Открыть карточку", callback_data=f"{ADM_CB_USER_PREFIX}:{telegram_user_id}")],
        ]
    )


def _portfolio_analysis_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="1 день", callback_data=f"{PA_CB_PREFIX}:1d"),
                InlineKeyboardButton(text="7 дней", callback_data=f"{PA_CB_PREFIX}:7d"),
            ],
            [
                InlineKeyboardButton(text="Месяц", callback_data=f"{PA_CB_PREFIX}:1m"),
                InlineKeyboardButton(text="Год", callback_data=f"{PA_CB_PREFIX}:1y"),
            ],
        ]
    )


def _analysis_horizon_label(horizon: str | None) -> str:
    mapping = {
        "1d": "1 день",
        "7d": "7 дней",
        "1m": "1 месяц",
        "1y": "1 год",
    }
    return mapping.get(str(horizon or "").strip().lower(), "выбранный срок")


def _track_background_task(task: asyncio.Task) -> None:
    BACKGROUND_TASKS.add(task)
    task.add_done_callback(BACKGROUND_TASKS.discard)


def _analysis_docx_path(telegram_user_id: int, horizon: str) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"portfolio_analysis_{telegram_user_id}_{horizon}_{timestamp}.docx"
    return Path(tempfile.gettempdir()) / filename


def _write_portfolio_analysis_docx(payload: dict, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    report = str(payload.get("report") or "").strip()
    horizon_label = str(payload.get("horizon_label") or "—")
    generated_at = _format_datetime(str(payload.get("generated_at_msk") or ""))
    model = str(payload.get("model") or "AI")

    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    document.add_heading("AI-анализ портфеля", level=1)
    meta = document.add_paragraph()
    meta.add_run("Горизонт: ").bold = True
    meta.add_run(horizon_label)
    meta.add_run("\nДата МСК: ").bold = True
    meta.add_run(generated_at)
    meta.add_run("\nМодель: ").bold = True
    meta.add_run(model)

    document.add_paragraph("")
    for block in report.split("\n"):
        line = _clean_docx_report_line(block)
        if not line:
            document.add_paragraph("")
            continue
        if line.startswith(("#", "ИТОГ:", "ФАКТОРЫ:", "ПО ПОЗИЦИЯМ:", "СЦЕНАРИЙ:", "РИСКИ И ТРИГГЕРЫ:")):
            document.add_heading(line.lstrip("# ").strip(), level=2)
        elif line.startswith(("1)", "1.", "2)", "2.", "3)", "3.", "4)", "4.", "5)", "5.", "6)", "6.")):
            document.add_heading(line, level=2)
        elif line.startswith(("-", "•")):
            document.add_paragraph(line.lstrip("-• ").strip(), style="List Bullet")
        else:
            document.add_paragraph(line)

    document.add_paragraph("")
    note = document.add_paragraph("Не является индивидуальной инвестиционной рекомендацией.")
    note.runs[0].italic = True
    document.save(path)


def _clean_docx_report_line(raw_line: str) -> str:
    line = str(raw_line or "").strip()
    line = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
    line = re.sub(r"__(.*?)__", r"\1", line)
    return line


async def _send_portfolio_analysis_document(
    *,
    bot: Bot,
    chat_id: int,
    telegram_user_id: int,
    horizon: str,
    horizon_label: str,
) -> None:
    path = _analysis_docx_path(telegram_user_id, horizon)
    try:
        payload = await service.analyze_portfolio(telegram_user_id, horizon=horizon)
        report = str(payload.get("report") or "").strip()
        if not report:
            await bot.send_message(chat_id=chat_id, text="AI вернул пустой отчет. Попробуйте повторить позже.")
            return

        await asyncio.to_thread(_write_portfolio_analysis_docx, payload, path)
        generated_at = _format_datetime(str(payload.get("generated_at_msk") or ""))
        caption = (
            "🤖 <b>AI-анализ портфеля готов</b>\n"
            f"Горизонт: <b>{html.escape(str(payload.get('horizon_label') or horizon_label))}</b>\n"
            f"Дата МСК: <b>{html.escape(generated_at)}</b>"
        )
        await bot.send_document(
            chat_id=chat_id,
            document=FSInputFile(str(path), filename=f"AI-анализ портфеля {horizon_label}.docx"),
            caption=caption,
        )
    except RuntimeError as exc:
        await bot.send_message(chat_id=chat_id, text=_portfolio_analysis_error_text(exc))
    except Exception:
        await bot.send_message(
            chat_id=chat_id,
            text="Не получилось сформировать Word-файл с отчетом. Попробуйте повторить анализ чуть позже.",
        )
    finally:
        try:
            path.unlink(missing_ok=True)
        except Exception:
            pass


def _calendar_keyboard(settings_payload: dict | None = None) -> InlineKeyboardMarkup:
    enabled = bool((settings_payload or {}).get("enabled", True))
    toggle_text = "Выключить уведомления" if enabled else "Включить уведомления"
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="Обновить", callback_data=CAL_CB_REFRESH),
                InlineKeyboardButton(text="Настройки", callback_data=CAL_CB_SETTINGS),
            ],
            [InlineKeyboardButton(text=toggle_text, callback_data=CAL_CB_TOGGLE)],
        ]
    )


def _calendar_settings_keyboard(settings_payload: dict) -> InlineKeyboardMarkup:
    current_days = _safe_int(settings_payload.get("days_before")) or 3
    enabled = bool(settings_payload.get("enabled", True))
    rows = [
        [
            InlineKeyboardButton(
                text=("✓ " if current_days == value else "") + f"{value} дн.",
                callback_data=f"{CAL_CB_DAYS_PREFIX}:{value}",
            )
            for value in (1, 3, 7)
        ],
        [
            InlineKeyboardButton(
                text=("✓ " if current_days == value else "") + f"{value} дн.",
                callback_data=f"{CAL_CB_DAYS_PREFIX}:{value}",
            )
            for value in (14, 30)
        ],
        [InlineKeyboardButton(text="Выключить" if enabled else "Включить", callback_data=CAL_CB_TOGGLE)],
        [InlineKeyboardButton(text="Назад к календарю", callback_data=CAL_CB_REFRESH)],
    ]
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _render_calendar_text(payload: dict, settings_payload: dict | None = None) -> str:
    items = payload.get("items") if isinstance(payload, dict) else []
    items = items if isinstance(items, list) else []
    settings_payload = settings_payload or {}
    enabled = bool(settings_payload.get("enabled", True))
    days_before = _safe_int(settings_payload.get("days_before")) or 3
    lines = [
        "🗓 <b>Календарь избранного</b>",
        f"Уведомления: <b>{'включены' if enabled else 'выключены'}</b>, за <b>{days_before} дн.</b>",
        "",
    ]
    if not items:
        lines.append("Событий по избранным инструментам пока не найдено.")
        lines.append("Добавьте акции, облигации или фонды в избранное, чтобы отслеживать выплаты.")
        return "\n".join(lines)

    for raw_item in items[:15]:
        if not isinstance(raw_item, dict):
            continue
        ticker = html.escape(str(raw_item.get("ticker") or raw_item.get("figi") or "—"))
        event_type = html.escape(str(raw_item.get("event_type_label") or raw_item.get("event_type") or "Событие"))
        event_date = html.escape(str(raw_item.get("event_date") or "—"))
        days_left = _safe_int(raw_item.get("days_left"))
        days_text = f"через {days_left} дн." if days_left is not None and days_left >= 0 else "дата прошла"
        amount = raw_item.get("amount")
        currency = raw_item.get("currency")
        amount_text = f" · {html.escape(str(amount))} {html.escape(str(currency or ''))}" if amount else ""
        description = raw_item.get("description")
        description_text = f"\n   {html.escape(str(description))}" if description else ""
        lines.append(
            f"• <b>{event_date}</b> · {event_type} · <b>{ticker}</b> ({days_text}){amount_text}{description_text}"
        )

    total = _safe_int(payload.get("total")) or len(items)
    if total > 15:
        lines.append(f"\nПоказаны ближайшие 15 из {total}.")
    errors = payload.get("errors")
    if isinstance(errors, list) and errors:
        lines.append(f"\nНе удалось загрузить часть событий: {len(errors)}.")
    return "\n".join(lines)


async def _send_calendar_message(message: Message) -> None:
    if not message.from_user:
        return
    try:
        settings_payload = await service.get_calendar_settings(message.from_user.id)
        payload = await service.get_calendar(message.from_user.id, days_ahead=180)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось загрузить календарь: {html.escape(str(exc))}"))
        return
    await _safe_telegram_call(
        message.answer(
            _render_calendar_text(payload, settings_payload),
            reply_markup=_calendar_keyboard(settings_payload),
        )
    )


async def _edit_calendar_message(callback: CallbackQuery) -> None:
    try:
        settings_payload = await service.get_calendar_settings(callback.from_user.id)
        payload = await service.get_calendar(callback.from_user.id, days_ahead=180)
    except RuntimeError as exc:
        await callback.answer(f"Не удалось загрузить календарь: {exc}", show_alert=True)
        return
    if callback.message:
        await _safe_edit_text(
            callback.message,
            _render_calendar_text(payload, settings_payload),
            reply_markup=_calendar_keyboard(settings_payload),
        )


async def _edit_calendar_settings(callback: CallbackQuery) -> None:
    try:
        settings_payload = await service.get_calendar_settings(callback.from_user.id)
    except RuntimeError as exc:
        await callback.answer(f"Не удалось загрузить настройки: {exc}", show_alert=True)
        return
    text = (
        "🔔 <b>Уведомления календаря</b>\n"
        f"Статус: <b>{'включены' if settings_payload.get('enabled') else 'выключены'}</b>\n"
        f"Напоминать за: <b>{settings_payload.get('days_before')} дн.</b>\n\n"
        "Уведомления отправляются по событиям из избранного: дивиденды, купоны и события облигаций."
    )
    if callback.message:
        await _safe_edit_text(callback.message, text, reply_markup=_calendar_settings_keyboard(settings_payload))


async def _notify_admins_access_request(message: Message, access_item: dict) -> None:
    if not message.from_user:
        return
    telegram_user_id = int(access_item.get("telegram_user_id") or message.from_user.id)
    name = _display_user_name(access_item)
    text = (
        "🆕 <b>Новая заявка на доступ к боту</b>\n"
        f"ID: <code>{telegram_user_id}</code>\n"
        f"Имя: <b>{html.escape(name)}</b>\n"
        f"Username: <b>{html.escape('@' + message.from_user.username if message.from_user.username else '—')}</b>\n"
        "Выберите действие:"
    )
    for admin_id in sorted(ADMIN_TELEGRAM_IDS):
        await _safe_telegram_call(
            message.bot.send_message(
                chat_id=admin_id,
                text=text,
                reply_markup=_admin_request_keyboard(telegram_user_id),
            )
        )


def _list_header(total: int, page: int, total_pages: int) -> str:
    return (
        "📊 <b>Инструменты T-Bank</b>\n"
        f"Всего: <b>{total}</b>\n"
        f"Страница: <b>{page}/{total_pages}</b>\n\n"
        "Кнопки сверху:\n"
        "• <b>Поиск</b> — найти по тикеру/FIGI/названию\n"
        "• <b>Избранное</b> — только ваши отмеченные бумаги\n"
        "• <b>Все</b> — полный список\n"
        "• <b>Акции/Облигации/Фонды</b> — фильтр по типу\n\n"
        "Выберите инструмент для подробной информации:"
    )


def _shares_mode_header(mode: str, query: str | None) -> str:
    if mode == "favorites":
        return "Режим: <b>Избранные</b>\n"
    if mode == "search":
        query_text = html.escape((query or "").strip() or "—")
        return f"Режим: <b>Поиск</b> · запрос: <code>{query_text}</code>\n"
    return "Режим: <b>Все инструменты</b>\n"


def _instrument_filter_label(value: str | None) -> str:
    mapping = {
        "share": "Акции",
        "bond": "Облигации",
        "etf": "Фонды",
        "all": "Все типы",
    }
    return mapping.get(str(value or "all").strip().lower(), "Все типы")


def _instrument_types_for_filter(value: str | None) -> list[str]:
    normalized = str(value or "all").strip().lower()
    if normalized in {"share", "bond", "etf"}:
        return [normalized]
    return ["share", "bond", "etf"]


def _order_submit_error_text(exc: RuntimeError) -> str:
    details = str(exc)
    if "30042" in details or "Not enough assets for a margin trade" in details:
        return (
            "Недостаточно активов для маржинальной сделки (код 30042).\n"
            "Пополните счет или уменьшите объем заявки."
        )
    if "30240" in details or "Confirmation required for specified instrument" in details:
        return (
            "Заявка требует подтверждения возможной непокрытой позиции (код 30240).\n"
            "Флаг confirm_margin_trade уже включён, повторите попытку через несколько секунд."
        )
    return f"Не удалось разместить заявку: {exc}"


def _portfolio_analysis_error_text(exc: RuntimeError) -> str:
    details = str(exc).strip()
    lowered = details.lower()
    if "unsupported analysis horizon" in lowered:
        return (
            "Не распознал срок прогноза.\n\n"
            "Откройте «🤖 Анализ портфеля» заново и выберите срок кнопкой."
        )
    if "insufficient balance" in lowered or "402" in lowered:
        return (
            "AI-сервис сейчас не может сформировать отчет: на балансе ProxyAPI недостаточно средств.\n\n"
            "Пополните баланс и повторите анализ."
        )
    if "api key" in lowered or "ключ" in lowered:
        return (
            "AI-анализ пока не настроен: сервер не видит ключ ProxyAPI.\n\n"
            "Проверьте переменную <code>AI_API_KEY</code> на сервере."
        )
    if (
        "больше времени" in lowered
        or "timeout" in lowered
        or "timed out" in lowered
        or "ai gateway" in lowered
        or "gateway" in lowered
    ):
        return (
            "AI-сервис не успел подготовить отчет.\n\n"
            "Портфель собран, но модель ответила слишком долго. Повторите попытку чуть позже."
        )
    if "t-bank" in lowered or "tbank" in lowered:
        return (
            "Не удалось получить данные портфеля из T-Bank.\n\n"
            "Проверьте профиль T-Bank и повторите попытку."
        )
    if "backend временно недоступен" in lowered:
        return (
            "Сервер бота сейчас не ответил на запрос анализа.\n\n"
            "Попробуйте еще раз через минуту."
        )
    return (
        "Не получилось сформировать AI-отчет.\n\n"
        "Попробуйте еще раз через минуту. Если ошибка повторится, я проверю логи сервера."
    )


def _shares_filter_items(
    items: list[ShareViewItem],
    *,
    mode: str,
    query: str | None,
    favorites: set[str],
    instrument_filter: str = "all",
) -> list[ShareViewItem]:
    allowed_types = set(_instrument_types_for_filter(instrument_filter))
    items = [item for item in items if item.instrument_type in allowed_types]
    if mode == "favorites":
        return [item for item in items if item.figi in favorites]
    if mode == "search":
        needle = (query or "").strip().upper()
        if not needle:
            return []
        return [
            item
            for item in items
            if needle in item.figi.upper()
            or needle in item.ticker.upper()
            or needle in item.name.upper()
            or needle in (item.isin or "").upper()
        ]
    return items




def _format_price(value: str | None, currency: str) -> str:
    if not value:
        return "—"
    try:
        dec = Decimal(value)
    except (InvalidOperation, ValueError):
        return value
    normalized = format(dec.normalize(), "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return f"{normalized} {currency}".strip()


def _format_percent(value: str | None) -> str:
    if not value:
        return "—"
    try:
        dec = Decimal(value)
    except (InvalidOperation, ValueError):
        return value
    rounded = dec.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    normalized = format(rounded.normalize(), "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    if normalized in ("-0", "-0.0", "-0.00"):
        normalized = "0"
    sign = "+" if rounded > 0 else ""
    return f"{sign}{normalized}%"


def _format_datetime(value: str | None) -> str:
    if not value:
        return "—"
    raw = value.strip()
    if not raw:
        return "—"
    try:
        parsed = datetime.fromisoformat(raw)
        return parsed.strftime("%d.%m.%Y %H:%M:%S")
    except ValueError:
        return raw




def _format_json_short(payload: dict, limit: int = 3000) -> str:
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    return html.escape(text[:limit])


def _money_from_quotation(raw: dict | None) -> Decimal | None:
    if not isinstance(raw, dict):
        return None
    try:
        units = int(raw.get("units", 0))
        nano = int(raw.get("nano", 0))
    except Exception:
        return None
    return Decimal(units) + (Decimal(nano) / Decimal(1_000_000_000))


def _format_money(raw: dict | None, fallback_currency: str = "RUB") -> str:
    dec = _money_from_quotation(raw)
    if dec is None:
        return "—"
    currency = fallback_currency
    if isinstance(raw, dict):
        cur = str(raw.get("currency", "")).strip().upper()
        if cur:
            currency = cur
    normalized = format(dec.normalize(), "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return f"{normalized} {currency}"



def _format_decimal_human(value: Decimal | None, decimals: int = 2) -> str:
    if value is None:
        return "—"
    fmt = f"{{:,.{max(0, decimals)}f}}"
    text = fmt.format(value).replace(",", " ")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text


def _decimal_from_any(raw: object) -> Decimal | None:
    if isinstance(raw, dict):
        return _money_from_quotation(raw)
    if raw is None:
        return None
    try:
        return Decimal(str(raw))
    except Exception:
        return None


def _extract_currency(raw: dict | None, fallback_currency: str = "RUB") -> str:
    if isinstance(raw, dict):
        cur = str(raw.get("currency", "")).strip().upper()
        if cur:
            return cur
    return fallback_currency


def _format_decimal_plain(value: Decimal | None) -> str:
    if value is None:
        return "—"
    normalized = format(value.normalize(), "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return normalized


def _to_clean_num_str(value: object, suffix: str = "") -> str:
    try:
        dec = Decimal(str(value))
    except Exception:
        return f"{value}{suffix}".strip()
    normalized = format(dec.normalize(), "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return f"{normalized}{suffix}".strip()


def _parse_interval_seconds(raw_text: str) -> int | None:
    text = (raw_text or "").strip().lower()
    if not text:
        return None

    if text.isdigit():
        return int(text)

    match = re.fullmatch(r"(\d+)\s*(с|сек|сек\.|секунд|sec|s)", text)
    if match:
        return int(match.group(1))

    return None


def _format_interval_seconds(raw_value: object) -> str:
    try:
        total = int(raw_value)
    except Exception:
        return "—"
    if total <= 0:
        return "—"

    hours, remainder = divmod(total, 3600)
    minutes, seconds = divmod(remainder, 60)
    parts: list[str] = []
    if hours:
        parts.append(f"{hours}ч")
    if minutes:
        parts.append(f"{minutes}м")
    if seconds or not parts:
        parts.append(f"{seconds}с")
    return " ".join(parts)


def _monitor_interval_keyboard() -> ReplyKeyboardMarkup:
    labels = [f"{value} сек" for value in MONITOR_INTERVAL_PRESETS_SECONDS]
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text=labels[0]), KeyboardButton(text=labels[1]), KeyboardButton(text=labels[2])],
            [KeyboardButton(text=labels[3]), KeyboardButton(text=labels[4]), KeyboardButton(text=labels[5])],
            [KeyboardButton(text=CANCEL_TEXT)],
        ],
        resize_keyboard=True,
        one_time_keyboard=True,
    )


def _extract_monitor_interval_seconds(item: dict) -> int | None:
    raw_seconds = item.get("interval_seconds")
    if raw_seconds is not None:
        try:
            return int(raw_seconds)
        except Exception:
            return None

    # Backward compatibility for old payloads.
    raw_minutes = item.get("interval_minutes")
    if raw_minutes is not None:
        try:
            return int(raw_minutes) * 60
        except Exception:
            return None
    return None


def _parse_monitor_id(value: object) -> int | None:
    if isinstance(value, int):
        return value if value > 0 else None
    if value is None:
        return None
    raw = str(value).strip()
    if not raw:
        return None
    if raw.isdigit():
        parsed = int(raw)
        return parsed if parsed > 0 else None
    try:
        parsed = int(Decimal(raw))
        return parsed if parsed > 0 else None
    except Exception:
        return None


def _parse_callback_user_id(data: str | None, prefix: str) -> int | None:
    raw = str(data or "")
    if not raw.startswith(f"{prefix}:"):
        return None
    value = raw.rsplit(":", maxsplit=1)[-1].strip()
    if not value.isdigit():
        return None
    user_id = int(value)
    return user_id if user_id > 0 else None


def _find_monitor_item(items: list[dict], monitor_id: int) -> dict | None:
    for item in items:
        if not isinstance(item, dict):
            continue
        if _parse_monitor_id(item.get("id")) == monitor_id:
            return item
    return None









async def _get_user_monitors_payload(telegram_user_id: int) -> dict:
    return await service.list_monitors(telegram_user_id)


async def _get_user_monitor_item(telegram_user_id: int, monitor_id: int) -> dict | None:
    payload = await _get_user_monitors_payload(telegram_user_id)
    items = payload.get("items", [])
    if not isinstance(items, list):
        return None
    return _find_monitor_item(items, monitor_id)


async def _show_monitor_list_message(message: Message) -> None:
    if not message.from_user:
        return
    try:
        payload = await _get_user_monitors_payload(message.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось загрузить мониторинги: {exc}"))
        return

    items = payload.get("items", [])
    if not isinstance(items, list) or not items:
        await _safe_telegram_call(
            message.answer(
                "🧭 <b>Мои мониторинги</b>\n"
                "Пока ничего нет.\n\n"
                "Создайте первый через «🔔 Новый мониторинг».",
                reply_markup=_menu_keyboard_for_user(message.from_user.id if message.from_user else None),
            )
        )
        return

    active_count = sum(1 for item in items if isinstance(item, dict) and bool(item.get("is_active")))

    await _safe_telegram_call(
        message.answer(
            "🧭 <b>Мои мониторинги</b>\n"
            f"Активных: <b>{active_count}</b> из <b>{len(items)}</b>\n"
            "🟢 — активен, ⚪ — выключен\n"
            "Выберите монитор:",
            reply_markup=_monitor_list_keyboard(items),
        )
    )
    return


MAIN_MENU_TEXTS = {
    MENU_SHARES,
    MENU_NEW_MONITOR,
    MENU_MONITORS,
    MENU_CALENDAR,
    MENU_PORTFOLIO,
    MENU_PORTFOLIO_ANALYSIS,
    MENU_OPERATIONS,
    MENU_PROFILE,
    MENU_ORDER,
    MENU_STOP,
    MENU_ACTIVE_ORDERS,
    MENU_ADMIN,
}

MAIN_MENU_TEXTS_NORMALIZED = {text.strip().casefold() for text in MAIN_MENU_TEXTS}


def _menu_text_equals(actual: str, expected: str) -> bool:
    return actual.strip().casefold() == expected.strip().casefold()

def _is_main_menu_text(text: str) -> bool:
    normalized = (text or "").strip().casefold()
    return bool(normalized and normalized in MAIN_MENU_TEXTS_NORMALIZED)


async def _edit_monitor_list_message(callback: CallbackQuery) -> None:
    if not callback.from_user or not callback.message:
        return
    try:
        payload = await _get_user_monitors_payload(callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Не удалось загрузить мониторинги: {exc}"))
        return

    items = payload.get("items", [])
    if not isinstance(items, list) or not items:
        await _safe_edit_text(
            callback.message,
            "🧭 <b>Мои мониторинги</b>\n"
            "Пока ничего нет.\n\n"
            "Создайте первый через «🔔 Новый мониторинг».",
        )
        return

    active_count = sum(1 for item in items if isinstance(item, dict) and bool(item.get("is_active")))

    await _safe_edit_text(
        callback.message,
        "🧭 <b>Мои мониторинги</b>\n"
        f"Активных: <b>{active_count}</b> из <b>{len(items)}</b>\n"
        "🟢 — активен, ⚪ — выключен\n"
        "Выберите монитор:",
        reply_markup=_monitor_list_keyboard(items),
    )
    return


def _order_type_map(label: str) -> str | None:
    normalized = (label or "").strip().casefold()
    mapping = {
        ORDER_TYPE_LIMIT_LABEL.casefold(): "ORDER_TYPE_LIMIT",
        ORDER_TYPE_MARKET_LABEL.casefold(): "ORDER_TYPE_MARKET",
        ORDER_TYPE_BESTPRICE_LABEL.casefold(): "ORDER_TYPE_BESTPRICE",
    }
    return mapping.get(normalized)


def _stop_type_map(label: str) -> str | None:
    normalized = (label or "").strip().casefold()
    mapping = {
        STOP_KIND_STOP_LOSS_LABEL.casefold(): "STOP_ORDER_TYPE_STOP_LOSS",
        STOP_KIND_STOP_LOSS_LIMIT_LABEL.casefold(): "STOP_ORDER_TYPE_STOP_LIMIT",
        STOP_KIND_TAKE_PROFIT_LABEL.casefold(): "STOP_ORDER_TYPE_TAKE_PROFIT",
    }
    return mapping.get(normalized)


def _direction_map(label: str) -> str | None:
    normalized = (label or "").strip().casefold()
    mapping = {
        DIRECTION_BUY_LABEL.casefold(): "ORDER_DIRECTION_BUY",
        DIRECTION_SELL_LABEL.casefold(): "ORDER_DIRECTION_SELL",
    }
    return mapping.get(normalized)


def _order_type_help_text() -> str:
    return (
        "🧾 <b>Новая заявка</b>\n"
        "Выберите тип:\n"
        f"• <b>{ORDER_TYPE_LIMIT_LABEL}</b> — вы задаете цену, исполнение только по ней или лучше.\n"
        f"• <b>{ORDER_TYPE_MARKET_LABEL}</b> — быстрое исполнение по текущей рыночной цене.\n"
        f"• <b>{ORDER_TYPE_BESTPRICE_LABEL}</b> — исполнение по лучшей доступной цене в стакане.\n\n"
        "Если важна точная цена — лимитная.\n"
        "Если важна скорость — рыночная/лучшая цена."
    )


def _stop_type_help_text() -> str:
    return (
        "🛑 <b>Новый стоп-приказ</b>\n"
        "Выберите тип:\n"
        f"• <b>{STOP_KIND_STOP_LOSS_LABEL}</b> — защита от убытка, при срабатывании отправляет рыночную заявку.\n"
        f"• <b>{STOP_KIND_STOP_LOSS_LIMIT_LABEL}</b> — защита от убытка, при срабатывании ставит лимитную заявку.\n"
        f"• <b>{STOP_KIND_TAKE_PROFIT_LABEL}</b> — фиксация прибыли по целевой цене."
    )


def _direction_help_text() -> str:
    return (
        "Выберите направление:\n"
        f"• <b>{DIRECTION_BUY_LABEL}</b> — открыть/увеличить длинную позицию.\n"
        f"• <b>{DIRECTION_SELL_LABEL}</b> — сократить/закрыть позицию (или открыть шорт, если доступно)."
    )


async def _load_items(state: FSMContext, force_refresh: bool = False) -> list[ShareViewItem]:
    data = await state.get_data()
    cached_all_items = data.get("shares_all_items")
    if isinstance(cached_all_items, list) and cached_all_items and not force_refresh:
        return [ShareViewItem(**item) for item in cached_all_items]

    all_items = await service.get_all_stored_shares(limit=5000)
    await state.update_data(shares_all_items=[asdict(item) for item in all_items])
    return all_items


async def _load_favorite_figies(state: FSMContext, telegram_user_id: int, force_refresh: bool = False) -> set[str]:
    data = await state.get_data()
    cached = data.get("favorite_figies")
    if isinstance(cached, list) and not force_refresh:
        return {str(figi) for figi in cached if str(figi).strip()}
    figies = await service.get_favorites(telegram_user_id)
    clean = [figi for figi in figies if figi]
    await state.update_data(favorite_figies=clean)
    return set(clean)


async def _load_filtered_items(
    state: FSMContext,
    *,
    telegram_user_id: int,
    force_refresh: bool = False,
) -> list[ShareViewItem]:
    data = await state.get_data()
    mode = data.get("shares_mode") if isinstance(data.get("shares_mode"), str) else "all"
    query = data.get("shares_query") if isinstance(data.get("shares_query"), str) else ""
    instrument_filter = data.get("instrument_filter") if isinstance(data.get("instrument_filter"), str) else "all"
    all_items = await _load_items(state=state, force_refresh=force_refresh)
    favorites = await _load_favorite_figies(state=state, telegram_user_id=telegram_user_id, force_refresh=force_refresh)
    filtered = _shares_filter_items(
        all_items,
        mode=mode,
        query=query,
        favorites=favorites,
        instrument_filter=instrument_filter,
    )
    await state.update_data(shares_filtered_items=[asdict(item) for item in filtered])
    return filtered


async def _load_page_items(
    state: FSMContext,
    *,
    page: int,
    telegram_user_id: int,
    force_refresh: bool = False,
) -> list[ShareViewItem]:
    data = await state.get_data()
    mode = data.get("shares_mode") if isinstance(data.get("shares_mode"), str) else "all"
    query = data.get("shares_query") if isinstance(data.get("shares_query"), str) else ""
    instrument_filter = data.get("instrument_filter") if isinstance(data.get("instrument_filter"), str) else "all"
    cached_page = data.get("cached_page")
    cached_items = data.get("shares_items_page")
    if isinstance(cached_page, int) and cached_page == page and isinstance(cached_items, list) and not force_refresh:
        return [ShareViewItem(**item) for item in cached_items]

    if mode == "all" and not (query or "").strip():
        items, total, safe_page, total_pages = await service.get_shares_page(
            page=page,
            page_size=PAGE_SIZE,
            instrument_types=_instrument_types_for_filter(instrument_filter),
        )
        await state.update_data(
            shares_items_page=[asdict(item) for item in items],
            cached_page=safe_page,
            total_items=total,
            total_pages=total_pages,
            current_page=safe_page,
        )
        return items

    filtered = await _load_filtered_items(
        state=state,
        telegram_user_id=telegram_user_id,
        force_refresh=force_refresh,
    )
    total = len(filtered)
    total_pages = max(1, (total + PAGE_SIZE - 1) // PAGE_SIZE)
    safe_page = min(max(page, 1), total_pages)
    start = (safe_page - 1) * PAGE_SIZE
    end = start + PAGE_SIZE
    items = filtered[start:end]
    await state.update_data(
        shares_items_page=[asdict(item) for item in items],
        cached_page=safe_page,
        total_items=total,
        total_pages=total_pages,
        current_page=safe_page,
    )
    return items


async def _ensure_user_credentials(message: Message) -> bool:
    if not message.from_user:
        return False
    try:
        status = await service.get_user_credentials_status(message.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось проверить профиль T-Bank: {exc}"))
        return False
    if bool(status.get("configured")):
        return True
    await _safe_telegram_call(
        message.answer(
            "Сначала настройте профиль T-Bank в разделе «🔐 Профиль T-Bank»."
        )
    )
    return False


async def _emit_monitor_events(message: Message) -> None:
    if not message.from_user:
        return
    try:
        result = await service.check_monitors(message.from_user.id)
    except RuntimeError:
        return
    events = result.get("events", [])
    if not isinstance(events, list):
        return
    for event in events:
        if not isinstance(event, dict):
            continue
        monitor_id = _parse_monitor_id(event.get("monitor_id"))
        keyboard = None
        if monitor_id:
            keyboard = InlineKeyboardMarkup(
                inline_keyboard=[
                    [
                        InlineKeyboardButton(
                            text="Удалить",
                            callback_data=f"{MON_CB_DEL_PREFIX}:{monitor_id}",
                        ),
                        InlineKeyboardButton(
                            text="Обновить базу",
                            callback_data=f"{MON_CB_REBASE_PREFIX}:{monitor_id}",
                        ),
                    ]
                ]
            )
        await _safe_telegram_call(
            message.answer(
                "🔔 <b>Сработал мониторинг цены</b>\n"
                f"Инструмент: <b>{html.escape(str(event.get('ticker') or event.get('figi') or '—'))}</b>\n"
                f"Монитор: <b>#{event.get('monitor_id')}</b>\n\n"
                f"Текущая цена: <b>{_to_clean_num_str(event.get('current_price'))} RUB</b>\n"
                f"Базовая цена: <b>{_to_clean_num_str(event.get('base_price'))} RUB</b>\n"
                f"Изменение: <b>{_to_clean_num_str(event.get('change_percent'), '%')}</b> / "
                f"<b>{_to_clean_num_str(event.get('change_rub'))} RUB</b>\n"
                f"Порог: <b>{_to_clean_num_str(event.get('threshold_percent'), '%')}</b> или "
                f"<b>{_to_clean_num_str(event.get('threshold_rub'))} RUB</b>\n"
                f"Время МСК: <b>{_format_datetime(str(event.get('triggered_at_msk') or ''))}</b>",
                reply_markup=keyboard,
            )
        )


async def _show_page(message: Message, state: FSMContext, page: int, force_refresh: bool = False) -> None:
    if not message.from_user:
        return
    try:
        page_items = await _load_page_items(
            state=state,
            page=page,
            telegram_user_id=message.from_user.id,
            force_refresh=force_refresh,
        )
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось получить список инструментов: {exc}"))
        return

    data = await state.get_data()
    safe_page = data.get("current_page") if isinstance(data.get("current_page"), int) else 1
    total = data.get("total_items") if isinstance(data.get("total_items"), int) else len(page_items)
    total_pages = data.get("total_pages") if isinstance(data.get("total_pages"), int) else 1
    mode = data.get("shares_mode") if isinstance(data.get("shares_mode"), str) else "all"
    query = data.get("shares_query") if isinstance(data.get("shares_query"), str) else ""
    instrument_filter = data.get("instrument_filter") if isinstance(data.get("instrument_filter"), str) else "all"
    await state.set_state(SharesBrowserStates.browsing_list)
    await state.update_data(current_page=safe_page)

    await _safe_telegram_call(
        message.answer(
            _shares_mode_header(mode=mode, query=query) + _list_header(total=total, page=safe_page, total_pages=total_pages),
            reply_markup=shares_list_keyboard(
                items=page_items,
                page=safe_page,
                total_pages=total_pages,
                mode=mode,
                has_query=bool((query or "").strip()),
                instrument_filter=instrument_filter,
            ),
        )
    )
    await _emit_monitor_events(message)


async def _edit_page(callback: CallbackQuery, state: FSMContext, page: int, force_refresh: bool = False) -> None:
    if not callback.from_user:
        return
    try:
        page_items = await _load_page_items(
            state=state,
            page=page,
            telegram_user_id=callback.from_user.id,
            force_refresh=force_refresh,
        )
    except RuntimeError as exc:
        if callback.message:
            await _safe_telegram_call(callback.message.answer(f"Не удалось получить список инструментов: {exc}"))
        return

    data = await state.get_data()
    safe_page = data.get("current_page") if isinstance(data.get("current_page"), int) else 1
    total = data.get("total_items") if isinstance(data.get("total_items"), int) else len(page_items)
    total_pages = data.get("total_pages") if isinstance(data.get("total_pages"), int) else 1
    mode = data.get("shares_mode") if isinstance(data.get("shares_mode"), str) else "all"
    query = data.get("shares_query") if isinstance(data.get("shares_query"), str) else ""
    instrument_filter = data.get("instrument_filter") if isinstance(data.get("instrument_filter"), str) else "all"
    await state.set_state(SharesBrowserStates.browsing_list)
    await state.update_data(current_page=safe_page)

    if callback.message:
        try:
            await _safe_telegram_call(
                callback.message.edit_text(
                    _shares_mode_header(mode=mode, query=query) + _list_header(total=total, page=safe_page, total_pages=total_pages),
                    reply_markup=shares_list_keyboard(
                        items=page_items,
                        page=safe_page,
                        total_pages=total_pages,
                        mode=mode,
                        has_query=bool((query or "").strip()),
                        instrument_filter=instrument_filter,
                    ),
                )
            )
        except TelegramBadRequest as exc:
            if "message is not modified" not in str(exc).lower():
                raise


@router.message(CommandStart())
async def cmd_start(message: Message) -> None:
    if not message.from_user:
        return

    telegram_user_id = message.from_user.id
    if is_admin_user(telegram_user_id):
        await _safe_telegram_call(
            message.answer(
                "Привет! Это MoexMonitoring.\n"
                "У вас роль администратора доступа. Управление доступом — в разделе «🛡 Админка».",
                reply_markup=_menu_keyboard_for_user(telegram_user_id),
            )
        )
        return

    try:
        access = await service.get_bot_access_status(telegram_user_id)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось проверить доступ: {exc}"))
        return

    if bool(access.get("is_allowed")):
        await _safe_telegram_call(
            message.answer(
                "Привет! Это MoexMonitoring.\nВыберите раздел в меню ниже.",
                reply_markup=_menu_keyboard_for_user(telegram_user_id),
            )
        )
        return

    try:
        requested = await service.request_bot_access(
            telegram_user_id,
            username=message.from_user.username,
            first_name=message.from_user.first_name,
            last_name=message.from_user.last_name,
        )
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось отправить заявку на доступ: {exc}"))
        return

    await _safe_telegram_call(
        message.answer(
            "⏳ Доступ к боту пока не выдан.\n"
            "Заявка отправлена администраторам, дождитесь одобрения.",
        )
    )
    await _notify_admins_access_request(message, requested)
    return


@router.message(StateFilter("*"), F.text.in_(MAIN_MENU_TEXTS))
async def menu_interrupt_router(message: Message, state: FSMContext) -> None:
    await _open_menu_section(message, state, message.text)


@router.message(Command("cancel"))
@router.message(StateFilter("*"), F.text == CANCEL_TEXT)
async def cancel_current_flow(message: Message, state: FSMContext) -> None:
    await state.clear()
    await _safe_telegram_call(
        message.answer(
            "Текущее действие отменено.",
            reply_markup=_menu_keyboard_for_user(message.from_user.id if message.from_user else None),
        )
    )


async def _admin_show_panel(message: Message, *, edit: bool = False) -> None:
    text = (
        "🛡 <b>Админка доступа</b>\n"
        "Здесь можно:\n"
        "• одобрять заявки на доступ\n"
        "• отзывать доступ у пользователей\n"
        "• смотреть список пользователей"
    )
    if edit:
        await _safe_edit_text(message, text, reply_markup=_admin_panel_keyboard())
    else:
        await _safe_telegram_call(message.answer(text, reply_markup=_admin_panel_keyboard()))


async def _admin_show_pending_list(message: Message, *, edit: bool = False) -> None:
    try:
        payload = await service.list_pending_bot_access(limit=100, offset=0)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось загрузить заявки: {exc}"))
        return
    items = payload.get("items", [])
    safe_items = items if isinstance(items, list) else []
    total = int(payload.get("total") or len(safe_items))
    text = _admin_users_list_text(safe_items, title="Заявки на доступ", total=total)
    keyboard = _admin_users_list_keyboard(safe_items)
    if edit:
        await _safe_edit_text(message, text, reply_markup=keyboard)
    else:
        await _safe_telegram_call(message.answer(text, reply_markup=keyboard))


async def _admin_show_users_list(message: Message, *, edit: bool = False) -> None:
    try:
        payload = await service.list_bot_access_users(
            limit=200,
            offset=0,
            exclude_telegram_user_ids=sorted(ADMIN_TELEGRAM_IDS),
        )
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось загрузить пользователей: {exc}"))
        return
    items = payload.get("items", [])
    safe_items = items if isinstance(items, list) else []
    total = int(payload.get("total") or len(safe_items))
    text = _admin_users_list_text(safe_items, title="Пользователи бота", total=total)
    keyboard = _admin_users_list_keyboard(safe_items)
    if edit:
        await _safe_edit_text(message, text, reply_markup=keyboard)
    else:
        await _safe_telegram_call(message.answer(text, reply_markup=keyboard))


async def _admin_show_user_details(message: Message, telegram_user_id: int, *, edit: bool = False) -> None:
    try:
        payload = await service.get_bot_access_status(telegram_user_id)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось загрузить пользователя: {exc}"))
        return
    text = _admin_user_detail_text(payload)
    keyboard = _admin_user_actions_keyboard(telegram_user_id, str(payload.get("status") or "none"))
    if edit:
        await _safe_edit_text(message, text, reply_markup=keyboard)
    else:
        await _safe_telegram_call(message.answer(text, reply_markup=keyboard))


@router.message(Command("admin"))
@router.message(F.text == MENU_ADMIN)
async def admin_panel_start(message: Message, state: FSMContext) -> None:
    if not message.from_user or not is_admin_user(message.from_user.id):
        await _safe_telegram_call(message.answer("Доступ к админке запрещен."))
        return
    await state.clear()
    await _safe_telegram_call(
        message.answer(
            "Открываю админку доступа.",
            reply_markup=_menu_keyboard_for_user(message.from_user.id),
        )
    )
    await _admin_show_panel(message, edit=False)


@router.callback_query(F.data == ADM_CB_PANEL)
async def admin_panel_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer())
    if not callback.from_user or not callback.message or not is_admin_user(callback.from_user.id):
        await _safe_telegram_call(callback.answer("Нет доступа", show_alert=True))
        return
    await _admin_show_panel(callback.message, edit=True)


@router.callback_query(F.data == ADM_CB_PENDING)
async def admin_pending_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer("Загружаю заявки..."))
    if not callback.from_user or not callback.message or not is_admin_user(callback.from_user.id):
        await _safe_telegram_call(callback.answer("Нет доступа", show_alert=True))
        return
    await _admin_show_pending_list(callback.message, edit=True)


@router.callback_query(F.data == ADM_CB_USERS)
async def admin_users_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer("Загружаю пользователей..."))
    if not callback.from_user or not callback.message or not is_admin_user(callback.from_user.id):
        await _safe_telegram_call(callback.answer("Нет доступа", show_alert=True))
        return
    await _admin_show_users_list(callback.message, edit=True)


@router.callback_query(F.data == ADM_CB_REFRESH)
async def admin_refresh_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer("Обновляю..."))
    if not callback.from_user or not callback.message or not is_admin_user(callback.from_user.id):
        await _safe_telegram_call(callback.answer("Нет доступа", show_alert=True))
        return
    await _admin_show_panel(callback.message, edit=True)


@router.callback_query(F.data.startswith(f"{ADM_CB_USER_PREFIX}:"))
async def admin_user_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer())
    if not callback.from_user or not callback.message or not is_admin_user(callback.from_user.id):
        await _safe_telegram_call(callback.answer("Нет доступа", show_alert=True))
        return
    telegram_user_id = _parse_callback_user_id(callback.data, ADM_CB_USER_PREFIX)
    if not telegram_user_id:
        await _safe_telegram_call(callback.answer("Некорректный ID", show_alert=True))
        return
    await _admin_show_user_details(callback.message, telegram_user_id, edit=True)


@router.callback_query(F.data.startswith(f"{ADM_CB_APPROVE_PREFIX}:"))
async def admin_approve_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer())
    if not callback.from_user or not callback.message or not is_admin_user(callback.from_user.id):
        await _safe_telegram_call(callback.answer("Нет доступа", show_alert=True))
        return
    telegram_user_id = _parse_callback_user_id(callback.data, ADM_CB_APPROVE_PREFIX)
    if not telegram_user_id:
        await _safe_telegram_call(callback.answer("Некорректный ID", show_alert=True))
        return
    try:
        payload = await service.approve_bot_access(
            telegram_user_id=telegram_user_id,
            admin_telegram_user_id=callback.from_user.id,
        )
    except RuntimeError as exc:
        await _safe_telegram_call(callback.answer(f"Не удалось выдать доступ: {exc}", show_alert=True))
        return
    await _safe_telegram_call(
        callback.bot.send_message(
            chat_id=telegram_user_id,
            text="✅ Доступ к боту выдан. Нажмите /start, чтобы открыть меню.",
        )
    )
    await _safe_edit_text(
        callback.message,
        _admin_user_detail_text(payload),
        reply_markup=_admin_user_actions_keyboard(telegram_user_id, str(payload.get("status") or "approved")),
    )


@router.callback_query(F.data.startswith(f"{ADM_CB_REVOKE_PREFIX}:"))
async def admin_revoke_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer())
    if not callback.from_user or not callback.message or not is_admin_user(callback.from_user.id):
        await _safe_telegram_call(callback.answer("Нет доступа", show_alert=True))
        return
    telegram_user_id = _parse_callback_user_id(callback.data, ADM_CB_REVOKE_PREFIX)
    if not telegram_user_id:
        await _safe_telegram_call(callback.answer("Некорректный ID", show_alert=True))
        return
    try:
        payload = await service.revoke_bot_access(
            telegram_user_id=telegram_user_id,
            admin_telegram_user_id=callback.from_user.id,
        )
    except RuntimeError as exc:
        await _safe_telegram_call(callback.answer(f"Не удалось отозвать доступ: {exc}", show_alert=True))
        return
    await _safe_telegram_call(
        callback.bot.send_message(
            chat_id=telegram_user_id,
            text="⛔ Доступ к боту отозван администратором.",
        )
    )
    await _safe_edit_text(
        callback.message,
        _admin_user_detail_text(payload),
        reply_markup=_admin_user_actions_keyboard(telegram_user_id, str(payload.get("status") or "revoked")),
    )


@router.message(Command("shares"))
@router.message(F.text == MENU_SHARES)
async def open_shares(message: Message, state: FSMContext) -> None:
    await state.update_data(
        shares_mode="all",
        shares_query="",
        instrument_filter="all",
        cached_page=None,
        shares_items_page=None,
    )
    await _show_page(message=message, state=state, page=1, force_refresh=True)


@router.message(Command("profile"))
@router.message(F.text == MENU_PROFILE)
async def profile_start(message: Message, state: FSMContext) -> None:
    await state.clear()
    if not message.from_user:
        return
    try:
        status = await service.get_user_credentials_status(message.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось загрузить профиль: {exc}"))
        return

    configured = bool(status.get("configured"))
    account_id = status.get("tbank_account_id")
    account_text = account_id if account_id else "—"
    status_text = "✅ Настроен" if configured else "⚠️ Не настроен"
    await _safe_telegram_call(
        message.answer(
            "🔐 <b>Профиль T-Bank</b>\n"
            f"Подключение: <b>{status_text}</b>\n"
            f"Telegram ID: <code>{message.from_user.id}</code>\n"
            f"Account ID: <code>{html.escape(str(account_text))}</code>\n\n"
            "Управление доступом:\n"
            "• <b>Обновить токен/account_id</b> — заменить реквизиты API.\n"
            "• <b>Проверить доступ</b> — быстрый тест подключения и прав.",
            reply_markup=_profile_keyboard(),
        )
    )
    return


@router.callback_query(F.data == PROF_CB_UPDATE)
async def profile_update_callback(callback: CallbackQuery, state: FSMContext) -> None:
    await _safe_telegram_call(callback.answer())
    await state.set_state(SharesBrowserStates.profile_set_token)
    if callback.message:
        prompt = await _safe_telegram_call(
            callback.message.answer(
                "Введите <b>T-Bank API token</b>:\n"
                "Сообщение будет удалено после обработки для безопасности."
            )
        )
        await state.update_data(profile_prompt_message_id=(prompt.message_id if prompt else None))


@router.callback_query(F.data == PROF_CB_CHECK)
async def profile_check_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer("Проверяю..."))
    if not callback.from_user or not callback.message:
        return
    try:
        status = await service.get_user_credentials_status(callback.from_user.id)
        if not bool(status.get("configured")):
            await _safe_telegram_call(
                callback.message.answer("Профиль не настроен. Сначала заполните токен и account_id.")
            )
            return
        portfolio = await service.get_portfolio(callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Проверка не пройдена: {exc}"))
        return
    details = portfolio.get("details", {})
    positions = details.get("positions", []) if isinstance(details, dict) else []
    await _safe_telegram_call(
        callback.message.answer(
            "Доступ к T-Bank подтвержден.\n"
            f"Позиций в портфеле: <b>{len(positions) if isinstance(positions, list) else 0}</b>"
        )
    )
    return


@router.message(SharesBrowserStates.profile_set_token)
async def profile_set_token(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    data = await state.get_data()
    await _safe_delete_message(message)
    await _safe_delete_by_id(message, data.get("profile_prompt_message_id"))
    token = message.text.strip()
    if len(token) < 10:
        await _safe_telegram_call(message.answer("Токен слишком короткий, попробуйте еще раз."))
        return
    await state.update_data(profile_token=token)
    await state.set_state(SharesBrowserStates.profile_set_account_id)
    prompt = await _safe_telegram_call(
        message.answer(
            "Введите <b>account_id</b> из T-Bank Invest:\n"
            "Это идентификатор брокерского счета, к которому будут отправляться заявки."
        )
    )
    await state.update_data(profile_prompt_message_id=(prompt.message_id if prompt else None))
    return


@router.message(SharesBrowserStates.profile_set_account_id)
async def profile_set_account_id(message: Message, state: FSMContext) -> None:
    if not message.text or not message.from_user:
        return
    data = await state.get_data()
    await _safe_delete_message(message)
    await _safe_delete_by_id(message, data.get("profile_prompt_message_id"))
    account_id = message.text.strip()
    if len(account_id) < 3:
        await _safe_telegram_call(message.answer("account_id выглядит слишком коротким, попробуйте еще раз."))
        return
    try:
        await service.upsert_user_credentials(
            telegram_user_id=message.from_user.id,
            tbank_token=data["profile_token"],
            tbank_account_id=account_id,
        )
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось сохранить профиль: {exc}"))
        return
    await state.clear()
    await _safe_telegram_call(
        message.answer(
            f"✅ Профиль сохранен.\nTelegram ID: <b>{message.from_user.id}</b>\nРеквизиты обновлены и скрыты из чата.",
            reply_markup=_menu_keyboard_for_user(message.from_user.id),
        )
    )
    return


@router.message(F.text == MENU_NEW_MONITOR)
async def monitor_start(message: Message, state: FSMContext) -> None:
    await state.set_state(SharesBrowserStates.monitoring_select_company)
    await _safe_telegram_call(
        message.answer(
            "🔔 <b>Новый мониторинг</b>\n"
            "Мониторинг отслеживает движение цены и присылает уведомления.\n\n"
            "Шаг 1/4: введите тикер или FIGI инструмента.\n"
            "Примеры: <code>SBER</code>, <code>BBG004730N88</code>."
        )
    )
    return


@router.message(SharesBrowserStates.monitoring_select_company)
async def monitor_set_company(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    if _is_main_menu_text(message.text):
        await _open_menu_section(message, state, message.text)
        return
    share = await service.find_share_by_ticker_or_figi(message.text)
    if share is None or not share.last_price:
        await _safe_telegram_call(message.answer("Инструмент не найден или по нему пока нет цены. Попробуйте другой тикер/FIGI."))
        return
    await state.update_data(
        monitor_figi=share.figi,
        monitor_ticker=share.ticker,
        monitor_name=share.name,
        monitor_base_price=share.last_price,
    )
    await state.set_state(SharesBrowserStates.monitoring_set_interval)
    await _safe_telegram_call(
        message.answer(
            "Шаг 2/4: задайте интервал проверки в секундах.\n"
            "Можно нажать готовую кнопку или ввести вручную (например: <code>45</code> или <code>45 сек</code>).\n"
            f"Допустимый диапазон: <b>{MONITOR_INTERVAL_MIN_SECONDS}..{MONITOR_INTERVAL_MAX_SECONDS}</b> сек.",
            reply_markup=_monitor_interval_keyboard(),
        )
    )
    return


@router.message(SharesBrowserStates.monitoring_set_interval)
async def monitor_set_interval(message: Message, state: FSMContext) -> None:
    if message.text and _is_main_menu_text(message.text):
        await _open_menu_section(message, state, message.text)
        return
    if not message.text:
        await _safe_telegram_call(
            message.answer(
                "Введите интервал в секундах (например: <code>30</code> или <code>30 сек</code>).",
                reply_markup=_monitor_interval_keyboard(),
            )
        )
        return
    interval_seconds = _parse_interval_seconds(message.text)
    if interval_seconds is None:
        await _safe_telegram_call(
            message.answer(
                "Не понял интервал. Введите секунды числом или с суффиксом «сек».\n"
                "Примеры: <code>15</code>, <code>30 сек</code>, <code>120</code>.",
                reply_markup=_monitor_interval_keyboard(),
            )
        )
        return
    if not (MONITOR_INTERVAL_MIN_SECONDS <= interval_seconds <= MONITOR_INTERVAL_MAX_SECONDS):
        await _safe_telegram_call(
            message.answer(
                f"Интервал должен быть в диапазоне {MONITOR_INTERVAL_MIN_SECONDS}..{MONITOR_INTERVAL_MAX_SECONDS} сек.",
                reply_markup=_monitor_interval_keyboard(),
            )
        )
        return
    await state.update_data(monitor_interval_seconds=interval_seconds)
    await state.set_state(SharesBrowserStates.monitoring_set_threshold_percent)
    await _safe_telegram_call(
        message.answer(
            "Шаг 3/4: введите порог в процентах.\n"
            "Уведомление придет, если цена изменится на этот процент от базовой.\n"
            "Пример: <code>5</code>."
        )
    )


@router.message(SharesBrowserStates.monitoring_set_threshold_percent)
async def monitor_set_percent(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    if _is_main_menu_text(message.text):
        await _open_menu_section(message, state, message.text)
        return
    try:
        Decimal(message.text.replace(",", ".").strip())
    except Exception:
        await _safe_telegram_call(message.answer("Введите корректное число."))
        return
    await state.update_data(monitor_percent=message.text.replace(",", ".").strip())
    await state.set_state(SharesBrowserStates.monitoring_set_threshold_rub)
    await _safe_telegram_call(
        message.answer(
            "Шаг 4/4: введите порог в RUB.\n"
            "Уведомление также придет, если изменение превысит эту сумму.\n"
            "Пример: <code>100</code>."
        )
    )


@router.message(SharesBrowserStates.monitoring_set_threshold_rub)
async def monitor_finish(message: Message, state: FSMContext) -> None:
    if not message.text or not message.from_user:
        return
    if _is_main_menu_text(message.text):
        await _open_menu_section(message, state, message.text)
        return
    try:
        Decimal(message.text.replace(",", ".").strip())
    except Exception:
        await _safe_telegram_call(message.answer("Введите корректное число."))
        return
    data = await state.get_data()
    try:
        created = await service.create_monitor(
            telegram_user_id=message.from_user.id,
            figi=data["monitor_figi"],
            ticker=data.get("monitor_ticker"),
            instrument_name=data.get("monitor_name"),
            interval_seconds=int(data["monitor_interval_seconds"]),
            threshold_percent=str(data["monitor_percent"]),
            threshold_rub=message.text.replace(",", ".").strip(),
            base_price=str(data["monitor_base_price"]),
        )
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось создать мониторинг: {exc}"))
        return

    await state.clear()
    details = created.get("details", {})
    interval_value = details.get("interval_seconds")
    if interval_value is None:
        interval_value = data.get("monitor_interval_seconds")
    await _safe_telegram_call(
        message.answer(
            "✅ Мониторинг создан\n"
            f"ID: <b>{details.get('id')}</b>\n"
            f"Инструмент: <b>{details.get('ticker') or details.get('figi')}</b>\n"
            f"Интервал: <b>{_format_interval_seconds(interval_value)}</b>\n"
            f"Порог: <b>{details.get('threshold_percent')}%</b> или <b>{details.get('threshold_rub')} RUB</b>",
            reply_markup=_menu_keyboard_for_user(message.from_user.id),
        )
    )
    return


@router.message(Command("monitors"))
@router.message(F.text == MENU_MONITORS)
async def list_monitors(message: Message) -> None:
    await _show_monitor_list_message(message)


@router.message(Command("monitor_check"))
async def monitor_check(message: Message) -> None:
    if not message.from_user:
        return
    await _emit_monitor_events(message)
    await _safe_telegram_call(message.answer("Проверка мониторингов выполнена."))


async def _monitor_state_change(message: Message, target_state: bool | None = None, delete: bool = False) -> None:
    if not message.from_user or not message.text:
        return
    parts = message.text.split(maxsplit=1)
    if len(parts) < 2 or not parts[1].strip().isdigit():
        await _safe_telegram_call(message.answer("Укажите ID: например /monitor_off 12"))
        return
    monitor_id = int(parts[1].strip())
    try:
        if delete:
            result = await service.delete_monitor(message.from_user.id, monitor_id)
            ok_text = f"✅ Монитор #{monitor_id} удален"
        else:
            result = await service.toggle_monitor(message.from_user.id, monitor_id, bool(target_state))
            ok_text = f"✅ Монитор #{monitor_id} {'включен' if target_state else 'выключен'}"
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Операция не выполнена: {exc}"))
        return
    if result.get("ok"):
        await _safe_telegram_call(message.answer(ok_text))
    return


@router.message(Command("monitor_on"))
async def monitor_on(message: Message) -> None:
    await _monitor_state_change(message, target_state=True)


@router.message(Command("monitor_off"))
async def monitor_off(message: Message) -> None:
    await _monitor_state_change(message, target_state=False)


@router.message(Command("monitor_del"))
async def monitor_del(message: Message) -> None:
    await _monitor_state_change(message, delete=True)


@router.callback_query(F.data == MON_CB_LIST)
async def monitors_list_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer())
    await _edit_monitor_list_message(callback)


@router.callback_query(F.data == MON_CB_REFRESH)
async def monitors_refresh_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer("\u041e\u0431\u043d\u043e\u0432\u043b\u044f\u044e \u0441\u043f\u0438\u0441\u043e\u043a..."))
    await _edit_monitor_list_message(callback)


@router.callback_query(F.data.startswith(f"{MON_CB_ITEM_PREFIX}:"))
async def monitor_item_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer())
    if not callback.from_user or not callback.message:
        return
    raw_id = callback.data.rsplit(":", maxsplit=1)[1]
    monitor_id = _parse_monitor_id(raw_id)
    if not monitor_id:
        await _safe_telegram_call(
            callback.answer("\u041d\u0435\u043a\u043e\u0440\u0440\u0435\u043a\u0442\u043d\u044b\u0439 ID \u043c\u043e\u043d\u0438\u0442\u043e\u0440\u0438\u043d\u0433\u0430", show_alert=True)
        )
        return
    try:
        payload = await _get_user_monitors_payload(callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Не удалось загрузить монитор: {exc}"))
        return
    items = payload.get("items", [])
    item = _find_monitor_item(items, monitor_id) if isinstance(items, list) else None
    if item is None:
        await _safe_telegram_call(callback.message.answer("  ."))
        return
    await _safe_edit_text(
        callback.message,
        _monitor_detail_text(item),
        reply_markup=_monitor_detail_keyboard(item),
    )


async def _monitor_action_callback(callback: CallbackQuery, action: str) -> None:
    if not callback.from_user or not callback.message:
        return
    raw_id = callback.data.rsplit(":", maxsplit=1)[1]
    monitor_id = _parse_monitor_id(raw_id)
    if not monitor_id:
        await _safe_telegram_call(
            callback.answer("\u041d\u0435\u043a\u043e\u0440\u0440\u0435\u043a\u0442\u043d\u044b\u0439 ID \u043c\u043e\u043d\u0438\u0442\u043e\u0440\u0438\u043d\u0433\u0430", show_alert=True)
        )
        return
    try:
        if action == "on":
            await service.toggle_monitor(callback.from_user.id, monitor_id, True)
        elif action == "off":
            await service.toggle_monitor(callback.from_user.id, monitor_id, False)
        elif action == "del":
            await service.delete_monitor(callback.from_user.id, monitor_id)
        elif action == "rebase":
            await service.rebase_monitor(callback.from_user.id, monitor_id)
        else:
            return
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Операция не выполнена: {exc}"))
        return

    if action == "del":
        await _safe_telegram_call(callback.answer("\u041c\u043e\u043d\u0438\u0442\u043e\u0440 \u0443\u0434\u0430\u043b\u0435\u043d"))
        await _edit_monitor_list_message(callback)
        return
    if action == "rebase":
        await _safe_telegram_call(callback.answer("\u0411\u0430\u0437\u043e\u0432\u0430\u044f \u0446\u0435\u043d\u0430 \u043e\u0431\u043d\u043e\u0432\u043b\u0435\u043d\u0430"))
    else:
        await _safe_telegram_call(callback.answer("\u041c\u043e\u043d\u0438\u0442\u043e\u0440 \u043e\u0431\u043d\u043e\u0432\u043b\u0435\u043d"))
    try:
        payload = await _get_user_monitors_payload(callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Не удалось обновить данные: {exc}"))
        return
    items = payload.get("items", [])
    item = _find_monitor_item(items, monitor_id) if isinstance(items, list) else None
    if item is None:
        await _edit_monitor_list_message(callback)
        return
    await _safe_edit_text(
        callback.message,
        _monitor_detail_text(item),
        reply_markup=_monitor_detail_keyboard(item),
    )


@router.callback_query(F.data.startswith(f"{MON_CB_ON_PREFIX}:"))
async def monitor_on_callback(callback: CallbackQuery) -> None:
    await _monitor_action_callback(callback, "on")


@router.callback_query(F.data.startswith(f"{MON_CB_OFF_PREFIX}:"))
async def monitor_off_callback(callback: CallbackQuery) -> None:
    await _monitor_action_callback(callback, "off")


@router.callback_query(F.data.startswith(f"{MON_CB_DEL_PREFIX}:"))
async def monitor_del_callback(callback: CallbackQuery) -> None:
    await _monitor_action_callback(callback, "del")


@router.callback_query(F.data.startswith(f"{MON_CB_REBASE_PREFIX}:"))
async def monitor_rebase_callback(callback: CallbackQuery) -> None:
    await _monitor_action_callback(callback, "rebase")


@router.callback_query(F.data.startswith(f"{MON_CB_EDIT_PREFIX}:"))
async def monitor_edit_callback(callback: CallbackQuery) -> None:
    if not callback.from_user or not callback.message:
        return
    raw_id = callback.data.rsplit(":", maxsplit=1)[1]
    monitor_id = _parse_monitor_id(raw_id)
    if not monitor_id:
        await _safe_telegram_call(
            callback.answer("\u041d\u0435\u043a\u043e\u0440\u0440\u0435\u043a\u0442\u043d\u044b\u0439 ID \u043c\u043e\u043d\u0438\u0442\u043e\u0440\u0438\u043d\u0433\u0430", show_alert=True)
        )
        return
    await _safe_telegram_call(callback.answer())
    try:
        item = await _get_user_monitor_item(callback.from_user.id, monitor_id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Не удалось загрузить монитор: {exc}"))
        return
    if item is None:
        await _safe_telegram_call(
            callback.answer("\u041c\u043e\u043d\u0438\u0442\u043e\u0440 \u043d\u0435 \u043d\u0430\u0439\u0434\u0435\u043d", show_alert=True)
        )
        return
    await _safe_edit_text(
        callback.message,
        _monitor_edit_menu_text(item),
        reply_markup=_monitor_edit_menu_keyboard(monitor_id),
    )


@router.callback_query(F.data.startswith(f"{MON_CB_EDIT_BACK_PREFIX}:"))
async def monitor_edit_back_callback(callback: CallbackQuery) -> None:
    if not callback.from_user or not callback.message:
        return
    raw_id = callback.data.rsplit(":", maxsplit=1)[1]
    monitor_id = _parse_monitor_id(raw_id)
    if not monitor_id:
        await _safe_telegram_call(
            callback.answer("\u041d\u0435\u043a\u043e\u0440\u0440\u0435\u043a\u0442\u043d\u044b\u0439 ID \u043c\u043e\u043d\u0438\u0442\u043e\u0440\u0438\u043d\u0433\u0430", show_alert=True)
        )
        return
    await _safe_telegram_call(callback.answer())
    try:
        item = await _get_user_monitor_item(callback.from_user.id, monitor_id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Не удалось загрузить монитор: {exc}"))
        return
    if item is None:
        await _safe_telegram_call(
            callback.answer("\u041c\u043e\u043d\u0438\u0442\u043e\u0440 \u043d\u0435 \u043d\u0430\u0439\u0434\u0435\u043d", show_alert=True)
        )
        return
    await _safe_edit_text(
        callback.message,
        _monitor_detail_text(item),
        reply_markup=_monitor_detail_keyboard(item),
    )


async def _start_monitor_edit_input(
    callback: CallbackQuery,
    state: FSMContext,
    *,
    target: str,
    prompt_text: str,
) -> None:
    if not callback.from_user or not callback.message:
        return
    raw_id = callback.data.rsplit(":", maxsplit=1)[1]
    monitor_id = _parse_monitor_id(raw_id)
    if not monitor_id:
        await _safe_telegram_call(
            callback.answer("\u041d\u0435\u043a\u043e\u0440\u0440\u0435\u043a\u0442\u043d\u044b\u0439 ID \u043c\u043e\u043d\u0438\u0442\u043e\u0440\u0438\u043d\u0433\u0430", show_alert=True)
        )
        return

    if target == "percent":
        await state.set_state(SharesBrowserStates.monitor_edit_threshold_percent)
    else:
        await state.set_state(SharesBrowserStates.monitor_edit_threshold_rub)

    prompt = await _safe_telegram_call(callback.message.answer(prompt_text))
    await state.update_data(
        edit_monitor_id=monitor_id,
        edit_target=target,
        edit_origin_message_id=callback.message.message_id,
        edit_prompt_message_id=(prompt.message_id if prompt else None),
    )
    await _safe_telegram_call(callback.answer())


@router.callback_query(F.data.startswith(f"{MON_CB_EDIT_PERCENT_PREFIX}:"))
async def monitor_edit_percent_callback(callback: CallbackQuery, state: FSMContext) -> None:
    await _start_monitor_edit_input(
        callback,
        state,
        target="percent",
        prompt_text="     ( 5  2.5):",
    )


@router.callback_query(F.data.startswith(f"{MON_CB_EDIT_RUB_PREFIX}:"))
async def monitor_edit_rub_callback(callback: CallbackQuery, state: FSMContext) -> None:
    await _start_monitor_edit_input(
        callback,
        state,
        target="rub",
        prompt_text="     ( 10):",
    )


async def _apply_monitor_edit_value(message: Message, state: FSMContext, target: str) -> None:
    if not message.text or not message.from_user:
        return
    data = await state.get_data()
    await _safe_delete_message(message)
    await _safe_delete_by_id(message, data.get("edit_prompt_message_id"))

    raw = message.text.replace(",", ".").strip()
    try:
        value = Decimal(raw)
    except Exception:
        prompt = await _safe_telegram_call(message.answer("  ."))
        await state.update_data(edit_prompt_message_id=(prompt.message_id if prompt else None))
        return

    if value < 0:
        prompt = await _safe_telegram_call(message.answer("   >= 0."))
        await state.update_data(edit_prompt_message_id=(prompt.message_id if prompt else None))
        return

    monitor_id = _parse_monitor_id(data.get("edit_monitor_id"))
    if not monitor_id:
        await state.clear()
        await _safe_telegram_call(
            message.answer(
                "   .",
                reply_markup=_menu_keyboard_for_user(message.from_user.id if message.from_user else None),
            )
        )
        return

    try:
        item = await _get_user_monitor_item(message.from_user.id, monitor_id)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось загрузить монитор: {exc}"))
        return

    if item is None:
        await state.clear()
        await _safe_telegram_call(
            message.answer(
                "  .",
                reply_markup=_menu_keyboard_for_user(message.from_user.id if message.from_user else None),
            )
        )
        return

    current_percent = str(item.get("threshold_percent"))
    current_rub = str(item.get("threshold_rub"))
    new_percent = str(value) if target == "percent" else current_percent
    new_rub = str(value) if target == "rub" else current_rub

    try:
        await service.update_monitor_thresholds(
            telegram_user_id=message.from_user.id,
            monitor_id=monitor_id,
            threshold_percent=new_percent,
            threshold_rub=new_rub,
        )
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось обновить пороги: {exc}"))
        return

    try:
        updated = await _get_user_monitor_item(message.from_user.id, monitor_id)
    except RuntimeError:
        updated = None

    origin_message_id = data.get("edit_origin_message_id")
    if updated is not None and isinstance(origin_message_id, int):
        await _safe_telegram_call(
            message.bot.edit_message_text(
                chat_id=message.chat.id,
                message_id=origin_message_id,
                text=_monitor_detail_text(updated),
                reply_markup=_monitor_detail_keyboard(updated),
            )
        )

    await state.clear()


@router.message(SharesBrowserStates.monitor_edit_threshold_percent)
async def monitor_edit_set_percent(message: Message, state: FSMContext) -> None:
    await _apply_monitor_edit_value(message, state, "percent")


@router.message(SharesBrowserStates.monitor_edit_threshold_rub)
async def monitor_edit_set_rub(message: Message, state: FSMContext) -> None:
    await _apply_monitor_edit_value(message, state, "rub")


@router.message(Command("calendar"))
@router.message(F.text == MENU_CALENDAR)
async def calendar_start(message: Message) -> None:
    if not message.from_user:
        return
    await _send_calendar_message(message)


@router.callback_query(F.data == CAL_CB_REFRESH)
async def calendar_refresh_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer("Обновляю календарь..."))
    await _edit_calendar_message(callback)


@router.callback_query(F.data == CAL_CB_SETTINGS)
async def calendar_settings_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer())
    await _edit_calendar_settings(callback)


@router.callback_query(F.data == CAL_CB_TOGGLE)
async def calendar_toggle_callback(callback: CallbackQuery) -> None:
    try:
        settings_payload = await service.get_calendar_settings(callback.from_user.id)
        updated = await service.update_calendar_settings(
            callback.from_user.id,
            enabled=not bool(settings_payload.get("enabled", True)),
            days_before=_safe_int(settings_payload.get("days_before")) or 3,
        )
    except RuntimeError as exc:
        await callback.answer(f"Не удалось обновить настройки: {exc}", show_alert=True)
        return
    await _safe_telegram_call(callback.answer("Настройки обновлены"))
    if callback.message:
        await _safe_edit_text(
            callback.message,
            (
                "🔔 <b>Уведомления календаря</b>\n"
                f"Статус: <b>{'включены' if updated.get('enabled') else 'выключены'}</b>\n"
                f"Напоминать за: <b>{updated.get('days_before')} дн.</b>"
            ),
            reply_markup=_calendar_settings_keyboard(updated),
        )


@router.callback_query(F.data.startswith(f"{CAL_CB_DAYS_PREFIX}:"))
async def calendar_days_callback(callback: CallbackQuery) -> None:
    raw_days = str(callback.data or "").rsplit(":", maxsplit=1)[-1]
    if not raw_days.isdigit():
        await callback.answer("Некорректное значение", show_alert=True)
        return
    days_before = int(raw_days)
    try:
        settings_payload = await service.get_calendar_settings(callback.from_user.id)
        updated = await service.update_calendar_settings(
            callback.from_user.id,
            enabled=bool(settings_payload.get("enabled", True)),
            days_before=days_before,
        )
    except RuntimeError as exc:
        await callback.answer(f"Не удалось обновить настройки: {exc}", show_alert=True)
        return
    await _safe_telegram_call(callback.answer(f"Буду напоминать за {days_before} дн."))
    if callback.message:
        await _safe_edit_text(
            callback.message,
            (
                "🔔 <b>Уведомления календаря</b>\n"
                f"Статус: <b>{'включены' if updated.get('enabled') else 'выключены'}</b>\n"
                f"Напоминать за: <b>{updated.get('days_before')} дн.</b>\n\n"
                "Уведомления отправляются по событиям из избранного."
            ),
            reply_markup=_calendar_settings_keyboard(updated),
        )


@router.message(F.text == MENU_PORTFOLIO)
async def show_portfolio(message: Message) -> None:
    if not message.from_user:
        return
    if not await _ensure_user_credentials(message):
        return
    try:
        payload = await service.get_portfolio(message.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось получить портфель: {exc}"))
        return
    await _safe_telegram_call(
        message.answer(_render_portfolio(payload.get("details", {})))
    )


@router.message(F.text == MENU_PORTFOLIO_ANALYSIS)
async def portfolio_analysis_start(message: Message) -> None:
    if not message.from_user:
        return
    if not await _ensure_user_credentials(message):
        return
    await _safe_telegram_call(
        message.answer(
            "🤖 <b>Анализ портфеля</b>\n"
            "Выберите горизонт прогноза. Я соберу текущий портфель T-Bank и подготовлю AI-отчет.",
            reply_markup=_portfolio_analysis_keyboard(),
        )
    )


@router.callback_query(F.data.startswith(f"{PA_CB_PREFIX}:"))
async def portfolio_analysis_callback(callback: CallbackQuery) -> None:
    if not callback.from_user or not callback.message:
        return
    horizon = str(callback.data or "").rsplit(":", maxsplit=1)[-1]
    await _safe_telegram_call(callback.answer("Готовлю анализ..."))
    horizon_label = _analysis_horizon_label(horizon)
    await _safe_telegram_call(
        callback.message.answer(
            "🤖 <b>Запустил анализ портфеля</b>\n"
            f"Горизонт: <b>{html.escape(horizon_label)}</b>\n\n"
            "Собираю позиции, структуру и доходность в фоне. Когда отчет будет готов, пришлю его Word-файлом в этот чат."
        )
    )
    task = asyncio.create_task(
        _send_portfolio_analysis_document(
            bot=callback.bot,
            chat_id=callback.message.chat.id,
            telegram_user_id=callback.from_user.id,
            horizon=horizon,
            horizon_label=horizon_label,
        )
    )
    _track_background_task(task)


@router.message(F.text == MENU_OPERATIONS)
async def operations_start(message: Message, state: FSMContext) -> None:
    if not message.from_user:
        return
    if not await _ensure_user_credentials(message):
        return
    await state.set_state(SharesBrowserStates.operations_set_days)
    await _safe_telegram_call(
        message.answer(
            "🕘 <b>Операции</b>\n"
            "За сколько дней показать историю?\n"
            "Введите число от <b>1</b> до <b>365</b> (например: <code>30</code>)."
        )
    )


@router.message(SharesBrowserStates.operations_set_days)
async def operations_show(message: Message, state: FSMContext) -> None:
    if not message.from_user or not message.text or not message.text.strip().isdigit():
        await _safe_telegram_call(message.answer("Введите целое число дней."))
        return
    days = int(message.text.strip())
    if days < 1 or days > 365:
        await _safe_telegram_call(message.answer("Число дней должно быть в диапазоне 1..365."))
        return
    try:
        payload = await service.get_operations(message.from_user.id, days=days)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось получить операции: {exc}"))
        return
    await state.clear()
    await _safe_telegram_call(
        message.answer(_render_operations(payload.get("details", {}), days=days))
    )


@router.message(F.text == MENU_ORDER)
async def order_start(message: Message, state: FSMContext) -> None:
    if not await _ensure_user_credentials(message):
        return
    await state.set_state(SharesBrowserStates.order_select_type)
    kb = ReplyKeyboardMarkup(
        keyboard=[
            [
                KeyboardButton(text=ORDER_TYPE_LIMIT_LABEL),
                KeyboardButton(text=ORDER_TYPE_MARKET_LABEL),
                KeyboardButton(text=ORDER_TYPE_BESTPRICE_LABEL),
            ],
            [KeyboardButton(text=CANCEL_TEXT)],
        ],
        resize_keyboard=True,
        one_time_keyboard=True,
    )
    await _safe_telegram_call(message.answer(_order_type_help_text(), reply_markup=kb))
    return


@router.message(SharesBrowserStates.order_select_type)
async def order_select_type(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    mapped = _order_type_map(message.text.strip())
    if not mapped:
        await _safe_telegram_call(
            message.answer(
                "Выберите тип заявки кнопкой на клавиатуре:\n"
                f"• <b>{ORDER_TYPE_LIMIT_LABEL}</b>\n"
                f"• <b>{ORDER_TYPE_MARKET_LABEL}</b>\n"
                f"• <b>{ORDER_TYPE_BESTPRICE_LABEL}</b>"
            )
        )
        return
    await state.update_data(order_type=mapped)
    await state.set_state(SharesBrowserStates.order_select_figi)
    await _safe_telegram_call(
        message.answer(
            "Введите тикер или FIGI инструмента (или Отмена).\n"
            "Примеры: <code>SBER</code>, <code>BBG004730N88</code>."
        )
    )


@router.message(SharesBrowserStates.order_select_figi)
async def order_select_figi(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    share = await service.find_share_by_ticker_or_figi(message.text)
    if share is None:
        await _safe_telegram_call(message.answer("Инструмент не найден. Попробуйте другой тикер/FIGI."))
        return
    await state.update_data(order_figi=share.figi, order_ticker=share.ticker)
    await state.set_state(SharesBrowserStates.order_select_direction)
    kb = ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text=DIRECTION_BUY_LABEL), KeyboardButton(text=DIRECTION_SELL_LABEL)],
            [KeyboardButton(text=CANCEL_TEXT)],
        ],
        resize_keyboard=True,
        one_time_keyboard=True,
    )
    await _safe_telegram_call(message.answer(_direction_help_text(), reply_markup=kb))


@router.message(SharesBrowserStates.order_select_direction)
async def order_select_direction(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    direction = _direction_map(message.text.strip())
    if not direction:
        await _safe_telegram_call(message.answer(_direction_help_text()))
        return
    await state.update_data(order_direction=direction)
    await state.set_state(SharesBrowserStates.order_set_quantity)
    await _safe_telegram_call(
        message.answer(
            "Введите количество <b>лотов</b> (или Отмена).\n"
            "Важно: вводится число лотов, не количество штук."
        )
    )


@router.message(SharesBrowserStates.order_set_quantity)
async def order_set_quantity(message: Message, state: FSMContext) -> None:
    if not message.text or not message.text.strip().isdigit():
        await _safe_telegram_call(message.answer("Введите целое число лотов."))
        return
    quantity = int(message.text.strip())
    if quantity <= 0:
        await _safe_telegram_call(message.answer("Количество должно быть больше 0."))
        return
    data = await state.get_data()
    await state.update_data(order_quantity=quantity)
    if data.get("order_type") == "ORDER_TYPE_LIMIT":
        await state.set_state(SharesBrowserStates.order_set_price)
        await _safe_telegram_call(
            message.answer(
                "Введите цену заявки (или Отмена).\n"
                "Это цена за 1 бумагу, по которой лимитная заявка будет стоять в стакане."
            )
        )
        return
    await _submit_order(message, state, None)


@router.message(SharesBrowserStates.order_set_price)
async def order_set_price(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    try:
        Decimal(message.text.replace(",", ".").strip())
    except Exception:
        await _safe_telegram_call(message.answer("Введите корректную цену."))
        return
    await _submit_order(message, state, message.text.replace(",", ".").strip())


async def _submit_order(message: Message, state: FSMContext, price: str | None) -> None:
    if not message.from_user:
        return
    data = await state.get_data()
    try:
        result = await service.create_order(
            telegram_user_id=message.from_user.id,
            figi=data["order_figi"],
            quantity_lots=int(data["order_quantity"]),
            direction=data["order_direction"],
            order_type=data["order_type"],
            price=price,
            confirm_margin_trade=True,
        )
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(_order_submit_error_text(exc)))
        return
    await state.clear()
    await _safe_telegram_call(
        message.answer("✅ Заявка отправлена\n<code>" + _format_json_short(result.get("details", {}), 1200) + "</code>")
    )
    return


@router.message(F.text == MENU_STOP)
async def stop_start(message: Message, state: FSMContext) -> None:
    if not await _ensure_user_credentials(message):
        return
    await state.set_state(SharesBrowserStates.stop_select_type)
    kb = ReplyKeyboardMarkup(
        keyboard=[
            [
                KeyboardButton(text=STOP_KIND_STOP_LOSS_LABEL),
                KeyboardButton(text=STOP_KIND_STOP_LOSS_LIMIT_LABEL),
                KeyboardButton(text=STOP_KIND_TAKE_PROFIT_LABEL),
            ],
            [KeyboardButton(text=CANCEL_TEXT)],
        ],
        resize_keyboard=True,
        one_time_keyboard=True,
    )
    await _safe_telegram_call(message.answer(_stop_type_help_text(), reply_markup=kb))
    return


@router.message(SharesBrowserStates.stop_select_type)
async def stop_select_type(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    mapped = _stop_type_map(message.text.strip())
    if not mapped:
        await _safe_telegram_call(message.answer(_stop_type_help_text()))
        return
    await state.update_data(stop_kind=message.text.strip(), stop_order_type=mapped)
    await state.set_state(SharesBrowserStates.stop_select_figi)
    await _safe_telegram_call(
        message.answer(
            "Введите тикер или FIGI инструмента (или Отмена).\n"
            "Примеры: <code>GAZP</code>, <code>BBG004730RP0</code>."
        )
    )


@router.message(SharesBrowserStates.stop_select_figi)
async def stop_select_figi(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    share = await service.find_share_by_ticker_or_figi(message.text)
    if share is None:
        await _safe_telegram_call(message.answer("Инструмент не найден. Попробуйте другой тикер/FIGI."))
        return
    await state.update_data(stop_figi=share.figi, stop_ticker=share.ticker)
    await state.set_state(SharesBrowserStates.stop_select_direction)
    kb = ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text=DIRECTION_BUY_LABEL), KeyboardButton(text=DIRECTION_SELL_LABEL)],
            [KeyboardButton(text=CANCEL_TEXT)],
        ],
        resize_keyboard=True,
        one_time_keyboard=True,
    )
    await _safe_telegram_call(
        message.answer(
            "Выберите направление для стоп-приказа:\n"
            f"• <b>{DIRECTION_BUY_LABEL}</b> — срабатывание на покупку.\n"
            f"• <b>{DIRECTION_SELL_LABEL}</b> — срабатывание на продажу.",
            reply_markup=kb,
        )
    )


@router.message(SharesBrowserStates.stop_select_direction)
async def stop_select_direction(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    direction = _direction_map(message.text.strip())
    if not direction:
        await _safe_telegram_call(
            message.answer(
                "Выберите направление кнопкой:\n"
                f"• <b>{DIRECTION_BUY_LABEL}</b>\n"
                f"• <b>{DIRECTION_SELL_LABEL}</b>"
            )
        )
        return
    await state.update_data(stop_direction=direction)
    await state.set_state(SharesBrowserStates.stop_set_quantity)
    await _safe_telegram_call(
        message.answer(
            "Введите количество <b>лотов</b> для стоп-приказа (или Отмена).\n"
            "Указывайте объем, который нужно защитить/зафиксировать."
        )
    )


@router.message(SharesBrowserStates.stop_set_quantity)
async def stop_set_quantity(message: Message, state: FSMContext) -> None:
    if not message.text or not message.text.strip().isdigit():
        await _safe_telegram_call(message.answer("Введите целое число лотов."))
        return
    quantity = int(message.text.strip())
    if quantity <= 0:
        await _safe_telegram_call(message.answer("Количество должно быть больше 0."))
        return
    await state.update_data(stop_quantity=quantity)
    await state.set_state(SharesBrowserStates.stop_set_stop_price)
    await _safe_telegram_call(
        message.answer(
            "Введите <b>стоп-цену</b> (или Отмена).\n"
            "При достижении этой цены активируется стоп-приказ."
        )
    )


@router.message(SharesBrowserStates.stop_set_stop_price)
async def stop_set_stop_price(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    try:
        Decimal(message.text.replace(",", ".").strip())
    except Exception:
        await _safe_telegram_call(message.answer("Введите корректную стоп-цену."))
        return
    await state.update_data(stop_price=message.text.replace(",", ".").strip())
    data = await state.get_data()
    if data.get("stop_kind") in {STOP_KIND_STOP_LOSS_LIMIT_LABEL, STOP_KIND_TAKE_PROFIT_LABEL}:
        await state.set_state(SharesBrowserStates.stop_set_limit_price)
        await _safe_telegram_call(
            message.answer(
                "Введите <b>лимитную цену</b> исполнения.\n"
                "Если хотите рыночное исполнение после срабатывания — введите <code>-</code>.\n"
                "Или нажмите Отмена."
            )
        )
        return
    await _submit_stop_order(message, state, None)


@router.message(SharesBrowserStates.stop_set_limit_price)
async def stop_set_limit_price(message: Message, state: FSMContext) -> None:
    if not message.text:
        return
    text = message.text.strip().replace(",", ".")
    if text != "-":
        try:
            Decimal(text)
        except Exception:
            await _safe_telegram_call(message.answer("Введите корректную цену или '-' ."))
            return
    await _submit_stop_order(message, state, None if text == "-" else text)


async def _submit_stop_order(message: Message, state: FSMContext, price: str | None) -> None:
    if not message.from_user:
        return
    data = await state.get_data()
    try:
        result = await service.create_stop_order(
            telegram_user_id=message.from_user.id,
            figi=data["stop_figi"],
            quantity_lots=int(data["stop_quantity"]),
            direction=data["stop_direction"],
            stop_order_type=data["stop_order_type"],
            stop_price=data["stop_price"],
            price=price,
        )
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось разместить стоп-приказ: {exc}"))
        return
    await state.clear()
    await _safe_telegram_call(
        message.answer("✅ Стоп-приказ отправлен\n<code>" + _format_json_short(result.get("details", {}), 1200) + "</code>")
    )
    return


@router.message(F.text == MENU_ACTIVE_ORDERS)
async def open_active_orders(message: Message) -> None:
    if not message.from_user:
        return
    if not await _ensure_user_credentials(message):
        return
    try:
        orders = await service.get_orders(message.from_user.id)
        stop_orders = await service.get_stop_orders(message.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(message.answer(f"Не удалось получить активные заявки: {exc}"))
        return
    orders_details = orders.get("details", {})
    stop_orders_details = stop_orders.get("details", {})
    active_orders, active_stops = _extract_active_orders_lists(orders_details, stop_orders_details)
    await _safe_telegram_call(
        message.answer(
            _render_active_orders(
                orders_details=orders_details,
                stop_orders_details=stop_orders_details,
            ),
            reply_markup=_active_orders_keyboard(active_orders, active_stops),
        )
    )


@router.callback_query(F.data == AO_CB_REFRESH)
async def active_orders_refresh_callback(callback: CallbackQuery) -> None:
    if not callback.from_user or not callback.message:
        return
    await _safe_telegram_call(callback.answer("Обновляю..."))
    try:
        orders = await service.get_orders(callback.from_user.id)
        stop_orders = await service.get_stop_orders(callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Не удалось получить активные заявки: {exc}"))
        return
    orders_details = orders.get("details", {})
    stop_orders_details = stop_orders.get("details", {})
    active_orders, active_stops = _extract_active_orders_lists(orders_details, stop_orders_details)
    await _safe_edit_text(
        callback.message,
        _render_active_orders(
            orders_details=orders_details,
            stop_orders_details=stop_orders_details,
        ),
        reply_markup=_active_orders_keyboard(active_orders, active_stops),
    )


@router.callback_query(F.data.startswith(f"{AO_CB_CANCEL_ORDER_PREFIX}:"))
async def active_orders_cancel_order_callback(callback: CallbackQuery) -> None:
    if not callback.from_user or not callback.message or not callback.data:
        return
    raw_index = callback.data.split(":", maxsplit=2)[-1]
    if not raw_index.isdigit() or int(raw_index) < 1:
        await _safe_telegram_call(callback.answer("Некорректный номер заявки", show_alert=True))
        return
    target_index = int(raw_index)

    await _safe_telegram_call(callback.answer("Отменяю биржевую заявку..."))
    try:
        orders = await service.get_orders(callback.from_user.id)
        stop_orders = await service.get_stop_orders(callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Не удалось получить активные заявки: {exc}"))
        return

    orders_details = orders.get("details", {})
    stop_orders_details = stop_orders.get("details", {})
    active_orders, active_stops = _extract_active_orders_lists(orders_details, stop_orders_details)
    if target_index > len(active_orders):
        await _safe_telegram_call(callback.answer("Эта заявка уже не активна", show_alert=True))
        await _safe_edit_text(
            callback.message,
            _render_active_orders(orders_details=orders_details, stop_orders_details=stop_orders_details),
            reply_markup=_active_orders_keyboard(active_orders, active_stops),
        )
        return

    selected = active_orders[target_index - 1]
    order_id = str(selected.get("orderId") or selected.get("order_id") or "").strip()
    if not order_id:
        await _safe_telegram_call(callback.answer("Не удалось определить ID заявки", show_alert=True))
        return

    try:
        await service.cancel_order(callback.from_user.id, order_id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.answer(f"Не удалось отменить заявку: {exc}", show_alert=True))
        return

    await _safe_telegram_call(callback.answer("Биржевая заявка отменена"))
    try:
        refreshed_orders = await service.get_orders(callback.from_user.id)
        refreshed_stop_orders = await service.get_stop_orders(callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Заявка отменена, но обновить список не удалось: {exc}"))
        return

    refreshed_orders_details = refreshed_orders.get("details", {})
    refreshed_stop_orders_details = refreshed_stop_orders.get("details", {})
    refreshed_active_orders, refreshed_active_stops = _extract_active_orders_lists(
        refreshed_orders_details,
        refreshed_stop_orders_details,
    )
    await _safe_edit_text(
        callback.message,
        _render_active_orders(
            orders_details=refreshed_orders_details,
            stop_orders_details=refreshed_stop_orders_details,
        ),
        reply_markup=_active_orders_keyboard(refreshed_active_orders, refreshed_active_stops),
    )


@router.callback_query(F.data.startswith(f"{AO_CB_CANCEL_STOP_PREFIX}:"))
async def active_orders_cancel_stop_callback(callback: CallbackQuery) -> None:
    if not callback.from_user or not callback.message or not callback.data:
        return
    raw_index = callback.data.split(":", maxsplit=2)[-1]
    if not raw_index.isdigit() or int(raw_index) < 1:
        await _safe_telegram_call(callback.answer("Некорректный номер стоп-приказа", show_alert=True))
        return
    target_index = int(raw_index)

    await _safe_telegram_call(callback.answer("Отменяю стоп-приказ..."))
    try:
        orders = await service.get_orders(callback.from_user.id)
        stop_orders = await service.get_stop_orders(callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Не удалось получить активные заявки: {exc}"))
        return

    orders_details = orders.get("details", {})
    stop_orders_details = stop_orders.get("details", {})
    active_orders, active_stops = _extract_active_orders_lists(orders_details, stop_orders_details)
    if target_index > len(active_stops):
        await _safe_telegram_call(callback.answer("Этот стоп-приказ уже не активен", show_alert=True))
        await _safe_edit_text(
            callback.message,
            _render_active_orders(orders_details=orders_details, stop_orders_details=stop_orders_details),
            reply_markup=_active_orders_keyboard(active_orders, active_stops),
        )
        return

    selected = active_stops[target_index - 1]
    stop_order_id = str(selected.get("stopOrderId") or selected.get("stop_order_id") or "").strip()
    if not stop_order_id:
        await _safe_telegram_call(callback.answer("Не удалось определить ID стоп-приказа", show_alert=True))
        return

    try:
        await service.cancel_stop_order(callback.from_user.id, stop_order_id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.answer(f"Не удалось отменить стоп-приказ: {exc}", show_alert=True))
        return

    await _safe_telegram_call(callback.answer("Стоп-приказ отменен"))
    try:
        refreshed_orders = await service.get_orders(callback.from_user.id)
        refreshed_stop_orders = await service.get_stop_orders(callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.message.answer(f"Стоп-приказ отменен, но обновить список не удалось: {exc}"))
        return

    refreshed_orders_details = refreshed_orders.get("details", {})
    refreshed_stop_orders_details = refreshed_stop_orders.get("details", {})
    refreshed_active_orders, refreshed_active_stops = _extract_active_orders_lists(
        refreshed_orders_details,
        refreshed_stop_orders_details,
    )
    await _safe_edit_text(
        callback.message,
        _render_active_orders(
            orders_details=refreshed_orders_details,
            stop_orders_details=refreshed_stop_orders_details,
        ),
        reply_markup=_active_orders_keyboard(refreshed_active_orders, refreshed_active_stops),
    )


@router.message(Command("cancel_order"))
async def cancel_order_cmd(message: Message) -> None:
    await _safe_telegram_call(
        message.answer(
            "Команды для отмены больше не используются.\n"
            "Откройте «📂 Активные заявки» и нажмите кнопку отмены у нужной заявки."
        )
    )


@router.message(Command("cancel_stop"))
async def cancel_stop_cmd(message: Message) -> None:
    await _safe_telegram_call(
        message.answer(
            "Команды для отмены больше не используются.\n"
            "Откройте «📂 Активные заявки» и нажмите кнопку отмены у нужного стоп-приказа."
        )
    )



@router.callback_query(F.data == CB_SEARCH)
async def shares_search_callback(callback: CallbackQuery, state: FSMContext) -> None:
    await _safe_telegram_call(callback.answer())
    await state.set_state(SharesBrowserStates.shares_search_query)
    if callback.message:
        await _safe_telegram_call(
            callback.message.answer(
                "\u0412\u0432\u0435\u0434\u0438\u0442\u0435 \u0437\u0430\u043f\u0440\u043e\u0441 (\u0442\u0438\u043a\u0435\u0440, \u043d\u0430\u0437\u0432\u0430\u043d\u0438\u0435, FIGI \u0438\u043b\u0438 ISIN):"
            )
        )


@router.message(SharesBrowserStates.shares_search_query)
async def shares_search_query_input(message: Message, state: FSMContext) -> None:
    if not message.text:
        await _safe_telegram_call(message.answer("\u0412\u0432\u0435\u0434\u0438\u0442\u0435 \u0442\u0435\u043a\u0441\u0442 \u0437\u0430\u043f\u0440\u043e\u0441\u0430."))
        return
    query = message.text.strip()
    if not query:
        await _safe_telegram_call(
            message.answer(
                "\u0417\u0430\u043f\u0440\u043e\u0441 \u043f\u0443\u0441\u0442\u043e\u0439. \u0412\u0432\u0435\u0434\u0438\u0442\u0435 \u0442\u0435\u043a\u0441\u0442 \u0437\u0430\u043f\u0440\u043e\u0441\u0430."
            )
        )
        return
    await state.update_data(
        shares_mode="search",
        shares_query=query,
        cached_page=None,
        shares_items_page=None,
    )
    await _show_page(message=message, state=state, page=1, force_refresh=False)


@router.callback_query(F.data == CB_MODE_ALL)
async def shares_all_mode_callback(callback: CallbackQuery, state: FSMContext) -> None:
    await _safe_telegram_call(callback.answer("Показываю все инструменты"))
    await state.update_data(
        shares_mode="all",
        shares_query="",
        cached_page=None,
        shares_items_page=None,
    )
    await _edit_page(callback=callback, state=state, page=1, force_refresh=False)


@router.callback_query(F.data == CB_MODE_FAVORITES)
async def shares_favorites_mode_callback(callback: CallbackQuery, state: FSMContext) -> None:
    await _safe_telegram_call(callback.answer("\u041f\u043e\u043a\u0430\u0437\u044b\u0432\u0430\u044e \u0438\u0437\u0431\u0440\u0430\u043d\u043d\u043e\u0435"))
    await state.update_data(
        shares_mode="favorites",
        shares_query="",
        cached_page=None,
        shares_items_page=None,
    )
    await _edit_page(callback=callback, state=state, page=1, force_refresh=False)


@router.callback_query(F.data.in_({CB_TYPE_ALL, CB_TYPE_SHARE, CB_TYPE_BOND, CB_TYPE_ETF}))
async def instruments_type_filter_callback(callback: CallbackQuery, state: FSMContext) -> None:
    mapping = {
        CB_TYPE_ALL: "all",
        CB_TYPE_SHARE: "share",
        CB_TYPE_BOND: "bond",
        CB_TYPE_ETF: "etf",
    }
    instrument_filter = mapping.get(callback.data or "", "all")
    await _safe_telegram_call(callback.answer(f"Фильтр: {_instrument_filter_label(instrument_filter)}"))
    await state.update_data(
        instrument_filter=instrument_filter,
        cached_page=None,
        shares_items_page=None,
    )
    await _edit_page(callback=callback, state=state, page=1, force_refresh=True)


@router.callback_query(F.data.startswith(f"{CB_FAVORITE_TOGGLE_PREFIX}:"))
async def toggle_favorite_callback(callback: CallbackQuery, state: FSMContext) -> None:
    if not callback.from_user or not callback.message:
        return
    data = await state.get_data()
    figi = str(data.get("details_figi") or "").strip()
    if not figi:
        await _safe_telegram_call(
            callback.answer("\u041d\u0435 \u0443\u0434\u0430\u043b\u043e\u0441\u044c \u043e\u043f\u0440\u0435\u0434\u0435\u043b\u0438\u0442\u044c \u0438\u043d\u0441\u0442\u0440\u0443\u043c\u0435\u043d\u0442", show_alert=True)
        )
        return

    raw_current = callback.data.split(":", maxsplit=1)[1]
    is_favorite = raw_current == "1"
    try:
        if is_favorite:
            await service.remove_favorite(callback.from_user.id, figi)
            toast = "\u0423\u0434\u0430\u043b\u0435\u043d\u043e \u0438\u0437 \u0438\u0437\u0431\u0440\u0430\u043d\u043d\u043e\u0433\u043e"
        else:
            await service.add_favorite(callback.from_user.id, figi)
            toast = "\u0414\u043e\u0431\u0430\u0432\u043b\u0435\u043d\u043e \u0432 \u0438\u0437\u0431\u0440\u0430\u043d\u043d\u043e\u0435"
    except RuntimeError as exc:
        await _safe_telegram_call(callback.answer(f"\u041e\u0448\u0438\u0431\u043a\u0430: {exc}", show_alert=True))
        return

    favorites = await _load_favorite_figies(
        state=state,
        telegram_user_id=callback.from_user.id,
        force_refresh=True,
    )
    item: ShareViewItem | None = None
    details_item = data.get("details_item")
    if isinstance(details_item, dict):
        try:
            parsed = ShareViewItem(**details_item)
            if parsed.figi == figi:
                item = parsed
        except TypeError:
            item = None
    if item is None:
        all_items = await _load_items(state=state, force_refresh=False)
        item = next((share for share in all_items if share.figi == figi), None)
    if item is not None:
        current_page = data.get("current_page") if isinstance(data.get("current_page"), int) else 1
        await state.update_data(details_item=asdict(item))
        await _safe_edit_text(
            callback.message,
            _share_details_text(item),
            reply_markup=share_details_keyboard(return_page=current_page, is_favorite=item.figi in favorites),
        )
    await _safe_telegram_call(callback.answer(toast))

@router.callback_query(F.data == "noop")
async def noop_callback(callback: CallbackQuery) -> None:
    await _safe_telegram_call(callback.answer())


@router.callback_query(F.data == CB_REFRESH)
async def refresh_list(callback: CallbackQuery, state: FSMContext) -> None:
    await _safe_telegram_call(callback.answer("\u041e\u0431\u043d\u043e\u0432\u043b\u044f\u044e \u0441\u043f\u0438\u0441\u043e\u043a..."))
    data = await state.get_data()
    raw_page = data.get("current_page")
    page = raw_page if isinstance(raw_page, int) else 1
    await _edit_page(callback=callback, state=state, page=page, force_refresh=True)


@router.callback_query(F.data == CB_REFRESH_DETAILS)
async def refresh_details(callback: CallbackQuery, state: FSMContext) -> None:
    await _safe_telegram_call(callback.answer("Обновляю..."))
    if not callback.from_user or not callback.message:
        return

    data = await state.get_data()
    figi = str(data.get("details_figi") or "").strip()
    if not figi:
        await _safe_telegram_call(callback.answer("Не удалось определить инструмент", show_alert=True))
        return

    try:
        item = await service.get_share_details_online(figi)
        favorites = await _load_favorite_figies(state=state, telegram_user_id=callback.from_user.id)
    except RuntimeError as exc:
        await _safe_telegram_call(callback.answer(f"Ошибка обновления: {exc}", show_alert=True))
        return

    current_page = data.get("current_page") if isinstance(data.get("current_page"), int) else 1
    await state.set_state(SharesBrowserStates.viewing_details)
    await state.update_data(current_page=current_page, details_figi=item.figi, details_item=asdict(item))
    await _safe_edit_text(
        callback.message,
        _share_details_text(item),
        reply_markup=share_details_keyboard(return_page=current_page, is_favorite=item.figi in favorites),
    )


@router.callback_query(F.data.startswith(f"{CB_PREFIX_PAGE}:"))
async def open_page(callback: CallbackQuery, state: FSMContext) -> None:
    await _safe_telegram_call(callback.answer())
    raw_page = callback.data.split(":", maxsplit=1)[1]
    page = int(raw_page) if raw_page.isdigit() else 1
    await _edit_page(callback=callback, state=state, page=page)


@router.callback_query(F.data.startswith(f"{CB_PREFIX_DETAILS}:"))
async def open_details(callback: CallbackQuery, state: FSMContext) -> None:
    await _safe_telegram_call(callback.answer("\u0417\u0430\u0433\u0440\u0443\u0436\u0430\u044e \u0434\u0435\u0442\u0430\u043b\u0438..."))
    if not callback.from_user:
        return
    figi = callback.data.split(":", maxsplit=1)[1]
    data = await state.get_data()
    current_page = data.get("current_page") if isinstance(data.get("current_page"), int) else 1
    try:
        favorites = await _load_favorite_figies(state=state, telegram_user_id=callback.from_user.id)
        item = await service.get_share_details_online(figi)
    except RuntimeError as exc:
        if callback.message:
            await _safe_telegram_call(callback.message.answer(f"Не удалось получить данные инструмента онлайн: {exc}"))
        return
    if item is None:
        if callback.message:
            await _safe_telegram_call(callback.message.answer("Инструмент не найден."))
        return
    await state.set_state(SharesBrowserStates.viewing_details)
    await state.update_data(current_page=current_page, details_figi=item.figi, details_item=asdict(item))
    if callback.message:
        try:
            await _safe_telegram_call(
                callback.message.edit_text(
                    _share_details_text(item),
                    reply_markup=share_details_keyboard(return_page=current_page, is_favorite=item.figi in favorites),
                )
            )
        except TelegramBadRequest as exc:
            if "message is not modified" not in str(exc).lower():
                raise

@router.callback_query(F.data.startswith(f"{CB_BACK_TO_LIST}:"))
async def back_to_list(callback: CallbackQuery, state: FSMContext) -> None:
    await _safe_telegram_call(callback.answer())
    raw_page = callback.data.split(":", maxsplit=1)[1]
    page = int(raw_page) if raw_page.isdigit() else 1
    await _edit_page(callback=callback, state=state, page=page)


@router.message(StateFilter("*"), F.text)
async def menu_fallback_router(message: Message, state: FSMContext) -> None:
    text = (message.text or "").strip()
    if _is_main_menu_text(text):
        await _open_menu_section(message, state, text)
    return




async def _open_menu_section(message: Message, state: FSMContext, text: str) -> None:
    await state.clear()
    normalized_text = (text or "").strip()
    if _menu_text_equals(normalized_text, MENU_SHARES):
        await open_shares(message, state)
        return
    if _menu_text_equals(normalized_text, MENU_NEW_MONITOR):
        await monitor_start(message, state)
        return
    if _menu_text_equals(normalized_text, MENU_MONITORS):
        await list_monitors(message)
        return
    if _menu_text_equals(normalized_text, MENU_CALENDAR):
        await calendar_start(message)
        return
    if _menu_text_equals(normalized_text, MENU_PORTFOLIO):
        await show_portfolio(message)
        return
    if _menu_text_equals(normalized_text, MENU_PORTFOLIO_ANALYSIS):
        await portfolio_analysis_start(message)
        return
    if _menu_text_equals(normalized_text, MENU_OPERATIONS):
        await operations_start(message, state)
        return
    if _menu_text_equals(normalized_text, MENU_PROFILE):
        await profile_start(message, state)
        return
    if _menu_text_equals(normalized_text, MENU_ORDER):
        await order_start(message, state)
        return
    if _menu_text_equals(normalized_text, MENU_STOP):
        await stop_start(message, state)
        return
    if _menu_text_equals(normalized_text, MENU_ACTIVE_ORDERS):
        await open_active_orders(message)
        return
    if _menu_text_equals(normalized_text, MENU_ADMIN):
        await admin_panel_start(message, state)
        return

    return


def _format_bool(value: bool | None) -> str:
    if value is True:
        return "Да"
    if value is False:
        return "Нет"
    return "—"


def _share_details_text_legacy(item: ShareViewItem) -> str:
    trading_text = "—"
    if item.trading_open is True:
        trading_text = "Идут"
    elif item.trading_open is False:
        trading_text = "\u0417\u0430\u043a\u0440\u044b\u0442\u0430"

    return (
        f"📊 <b>{item.name}</b>\n"
        f"Тип: <b>{_instrument_filter_label(item.instrument_type)}</b>\n"
        f"Тикер: <b>{item.ticker}</b>\n"
        f"FIGI: <code>{item.figi}</code>\n\n"
        f"Цена (последняя): <b>{_format_price(item.last_price, item.currency)}</b>\n"
        f"Время цены (МСК): <b>{_format_datetime(item.last_price_captured_at_msk)}</b>\n"
        f"Открытие (день): <b>{_format_price(item.day_open_price, item.currency)}</b>\n"
        f"Закрытие (пред. день): <b>{_format_price(item.day_close_price, item.currency)}</b>\n"
        f"Изменение за день: <b>{_format_percent(item.day_change_percent)}</b>\n"
        f"Изменение за год: <b>{_format_percent(item.year_change_percent)}</b>\n"
        f"Валюта: <b>{item.currency}</b>\n"
        f"Биржа: <b>{item.exchange}</b>\n"
        f"Страна риска: <b>{item.country_of_risk}</b>\n"
        f"Лот: <b>{item.lot if item.lot is not None else '—'}</b>\n"
        f"Номинал: <b>{item.nominal or '—'}</b>\n"
        f"ISIN: <code>{item.isin or '—'}</code>\n"
        f"Класс: <b>{item.class_code or '—'}</b>\n\n"
        f"Покупка доступна: <b>{_format_bool(item.buy_available)}</b>\n"
        f"Продажа доступна: <b>{_format_bool(item.sell_available)}</b>\n"
        f"API-торговля доступна: <b>{_format_bool(item.api_trade_available)}</b>\n"
        f"Шорт доступен: <b>{_format_bool(item.short_enabled)}</b>\n\n"
        f"🕘 <b>Торговая сессия</b>\n"
        f"Статус: <b>{trading_text}</b>\n"
        f"Открытие: <b>{_format_datetime(item.session_opened_at_msk)}</b>\n"
        f"Закрытие: <b>{_format_datetime(item.session_closed_at_msk)}</b>\n"
        f"След. открытие: <b>{_format_datetime(item.next_session_opened_at_msk)}</b>\n"
        f"След. закрытие: <b>{_format_datetime(item.next_session_closed_at_msk)}</b>"
    )


def _format_bool_legacy(value: bool | None) -> str:
    if value is True:
        return "Да"
    if value is False:
        return "Нет"
    return "—"


_TRADING_STATUS_LABELS: dict[str, str] = {
    "UNSPECIFIED": "Не определен",
    "NOT_AVAILABLE_FOR_TRADING": "Недоступен для торгов",
    "OPENING_PERIOD": "Период открытия",
    "CLOSING_PERIOD": "Период закрытия",
    "BREAK_IN_TRADING": "Перерыв в торгах",
    "NORMAL_TRADING": "Нормальная торговля",
    "CLOSING_AUCTION": "Аукцион закрытия",
    "DARK_POOL_AUCTION": "Dark pool",
    "DISCRETE_AUCTION": "Дискретный аукцион",
    "OPENING_AUCTION_PERIOD": "Аукцион открытия",
    "TRADING_AT_CLOSING_AUCTION_PRICE": "Торги по цене аукциона закрытия",
    "SESSION_ASSIGNED": "Сессия назначена",
    "SESSION_CLOSE": "Сессия закрыта",
    "SESSION_OPEN": "Сессия открыта",
    "DEALER_NORMAL_TRADING": "Внутренняя ликвидность брокера",
    "DEALER_BREAK_IN_TRADING": "Перерыв (внутренняя ликвидность)",
    "DEALER_NOT_AVAILABLE_FOR_TRADING": "Недоступно (внутренняя ликвидность)",
}


def _normalize_trading_status(value: str | None) -> str | None:
    raw = (value or "").strip().upper()
    if not raw:
        return None
    prefix = "SECURITY_TRADING_STATUS_"
    if raw.startswith(prefix):
        return raw.replace(prefix, "", 1)
    return raw


def _format_trading_status(value: str | None) -> str:
    normalized = _normalize_trading_status(value)
    if not normalized:
        return "—"
    return _TRADING_STATUS_LABELS.get(normalized, normalized)


def _can_buy_now(item: ShareViewItem) -> bool | None:
    if item.buy_available is False:
        return False
    if item.api_trade_available is False:
        return False
    if item.limit_order_available is False and item.market_order_available is False:
        return False
    if item.limit_order_available is True or item.market_order_available is True:
        return True
    return None


def _buy_now_reason(item: ShareViewItem) -> str:
    if item.buy_available is False:
        return "Инструмент недоступен для покупки (buyAvailableFlag=false)."
    if item.api_trade_available is False:
        return "Торговля через API недоступна (apiTradeAvailableFlag=false)."
    if item.limit_order_available is False and item.market_order_available is False:
        status = _format_trading_status(item.trading_status)
        return f"По текущему торговому статусу нельзя выставить заявку сейчас ({status})."
    if item.limit_order_available is True and item.market_order_available is False:
        return "Сейчас доступна только лимитная заявка."
    if item.market_order_available is True and item.limit_order_available is False:
        return "Сейчас доступна только рыночная заявка."
    if _can_buy_now(item) is True:
        return "Покупка сейчас доступна."
    return "Недостаточно данных от API для точного ответа."


def _share_details_text_legacy(item: ShareViewItem) -> str:
    trading_text = "—"
    if item.trading_open is True:
        trading_text = "Идут"
    elif item.trading_open is False:
        trading_text = "Закрыта"

    can_buy_now_text = _format_bool(_can_buy_now(item))
    buy_reason_text = _buy_now_reason(item)
    instrument_trading_status_text = _format_trading_status(item.trading_status)

    return (
        f"📊 <b>{item.name}</b>\n"
        f"Тикер: <b>{item.ticker}</b>\n"
        f"FIGI: <code>{item.figi}</code>\n\n"
        f"Цена (последняя): <b>{_format_price(item.last_price, item.currency)}</b>\n"
        f"Время цены (МСК): <b>{_format_datetime(item.last_price_captured_at_msk)}</b>\n"
        f"Открытие (день): <b>{_format_price(item.day_open_price, item.currency)}</b>\n"
        f"Закрытие (пред. день): <b>{_format_price(item.day_close_price, item.currency)}</b>\n"
        f"Изменение за день: <b>{_format_percent(item.day_change_percent)}</b>\n"
        f"Изменение за год: <b>{_format_percent(item.year_change_percent)}</b>\n"
        f"Валюта: <b>{item.currency}</b>\n"
        f"Биржа: <b>{item.exchange}</b>\n"
        f"Страна риска: <b>{item.country_of_risk}</b>\n"
        f"Лот: <b>{item.lot if item.lot is not None else '—'}</b>\n"
        f"Номинал: <b>{item.nominal or '—'}</b>\n"
        f"ISIN: <code>{item.isin or '—'}</code>\n"
        f"Класс: <b>{item.class_code or '—'}</b>\n\n"
        f"Покупка доступна: <b>{_format_bool(item.buy_available)}</b>\n"
        f"Продажа доступна: <b>{_format_bool(item.sell_available)}</b>\n"
        f"API-торговля доступна: <b>{_format_bool(item.api_trade_available)}</b>\n"
        f"Шорт доступен: <b>{_format_bool(item.short_enabled)}</b>\n\n"
        f"Статус торгов по бумаге: <b>{instrument_trading_status_text}</b>\n"
        f"Лимитная заявка сейчас: <b>{_format_bool(item.limit_order_available)}</b>\n"
        f"Рыночная заявка сейчас: <b>{_format_bool(item.market_order_available)}</b>\n"
        f"Можно купить сейчас: <b>{can_buy_now_text}</b>\n"
        f"Почему: <b>{buy_reason_text}</b>\n\n"
        f"🕘 <b>Торговая сессия</b>\n"
        f"Статус: <b>{trading_text}</b>\n"
        f"Открытие: <b>{_format_datetime(item.session_opened_at_msk)}</b>\n"
        f"Закрытие: <b>{_format_datetime(item.session_closed_at_msk)}</b>\n"
        f"След. открытие: <b>{_format_datetime(item.next_session_opened_at_msk)}</b>\n"
        f"След. закрытие: <b>{_format_datetime(item.next_session_closed_at_msk)}</b>"
    )


_share_details_text = _share_details_text_legacy


def _portfolio_type_label(raw_type: object) -> str:
    value = str(raw_type or "").strip().lower()
    labels = {
        "share": "Акция",
        "bond": "Облигация",
        "etf": "Фонд",
        "currency": "Валюта",
        "futures": "Фьючерс",
        "future": "Фьючерс",
        "option": "Опцион",
        "sp": "Структурный продукт",
        "clearing_certificate": "Клиринговый сертификат",
    }
    return labels.get(value, value.upper() if value else "Инструмент")


def _pl_badge(value: Decimal | None) -> str:
    if value is None or value == 0:
        return "⚪"
    return "🟢" if value > 0 else "🔴"


def _format_money_human(raw: dict | None, fallback_currency: str = "RUB", decimals: int = 2) -> str:
    dec = _money_from_quotation(raw)
    if dec is None:
        return "—"
    currency = _extract_currency(raw, fallback_currency)
    return f"{_format_decimal_human(dec, decimals=decimals)} {currency}"


def _format_signed_money(value: Decimal | None, currency: str = "RUB") -> str:
    if value is None:
        return "—"
    sign = "+" if value > 0 else ""
    return f"{sign}{_format_decimal_human(value)} {currency}"


def _format_signed_percent(value: Decimal | None) -> str:
    if value is None:
        return "—"
    sign = "+" if value > 0 else ""
    return f"{sign}{_format_decimal_human(value, decimals=2)}%"


def _portfolio_total_from_details(details: dict) -> Decimal:
    total_from_api = _money_from_quotation(details.get("totalAmountPortfolio"))
    if total_from_api is not None:
        return total_from_api
    return sum(
        (
            _money_from_quotation(details.get(key)) or Decimal("0")
            for key in (
                "totalAmountShares",
                "totalAmountBonds",
                "totalAmountEtf",
                "totalAmountCurrencies",
                "totalAmountFutures",
                "totalAmountOptions",
            )
        ),
        Decimal("0"),
    )


def _format_allocation_line(label: str, amount: Decimal | None, total: Decimal, currency: str) -> str:
    value = amount or Decimal("0")
    percent = "—"
    if total > 0:
        percent = _format_decimal_human((value / total * Decimal("100")).quantize(Decimal("0.01")), decimals=2)
    return f"• {label}: <b>{_format_decimal_human(value)} {currency}</b> · {percent}%"


def _position_market_value(pos: dict) -> Decimal | None:
    quantity = _money_from_quotation(pos.get("quantity"))
    current_price = _money_from_quotation(pos.get("currentPrice"))
    if quantity is None or current_price is None:
        return None
    return quantity * current_price


def _position_yield_percent(pos: dict) -> Decimal | None:
    avg_price = _money_from_quotation(pos.get("averagePositionPrice"))
    current_price = _money_from_quotation(pos.get("currentPrice"))
    if avg_price is None or current_price is None or avg_price == 0:
        return None
    return ((current_price - avg_price) / avg_price * Decimal("100")).quantize(Decimal("0.01"))


def _render_portfolio(details: dict) -> str:
    if not isinstance(details, dict):
        return "💼 <b>Портфель T-Bank</b>\nДанные портфеля пока недоступны."

    total_shares = _money_from_quotation(details.get("totalAmountShares"))
    total_bonds = _money_from_quotation(details.get("totalAmountBonds"))
    total_etf = _money_from_quotation(details.get("totalAmountEtf"))
    total_curr = _money_from_quotation(details.get("totalAmountCurrencies"))
    total_futures = _money_from_quotation(details.get("totalAmountFutures"))
    total_options = _money_from_quotation(details.get("totalAmountOptions"))
    expected_yield = _money_from_quotation(details.get("expectedYield"))
    daily_yield = _money_from_quotation(details.get("dailyYield"))
    currency = _extract_currency(details.get("totalAmountPortfolio"), _extract_currency(details.get("totalAmountShares"), "RUB"))
    total_all = _portfolio_total_from_details(details)
    yield_percent = None
    if expected_yield is not None and total_all - expected_yield != 0:
        yield_percent = (expected_yield / (total_all - expected_yield) * Decimal("100")).quantize(Decimal("0.01"))

    lines = [
        "💼 <b>Портфель T-Bank</b>",
        f"Итого: <b>{_format_decimal_human(total_all)} {currency}</b>",
        f"{_pl_badge(expected_yield)} Прибыль/убыток: <b>{_format_signed_money(expected_yield, currency)}</b> · {_format_signed_percent(yield_percent)}",
    ]
    if daily_yield is not None:
        lines.append(f"{_pl_badge(daily_yield)} За день: <b>{_format_signed_money(daily_yield, currency)}</b>")
    lines.extend([
        "",
        "📦 <b>Структура</b>",
        _format_allocation_line("Акции", total_shares, total_all, currency),
        _format_allocation_line("Облигации", total_bonds, total_all, currency),
        _format_allocation_line("Фонды", total_etf, total_all, currency),
        _format_allocation_line("Валюты", total_curr, total_all, currency),
        _format_allocation_line("Фьючерсы", total_futures, total_all, currency),
    ])
    if total_options is not None and total_options != 0:
        lines.append(_format_allocation_line("Опционы", total_options, total_all, currency))

    positions = details.get("positions")
    if not isinstance(positions, list) or not positions:
        lines.append("")
        lines.append("🧾 <b>Позиции</b>")
        lines.append("Открытых позиций пока нет.")
        return "\n".join(lines)

    valid_positions = [pos for pos in positions if isinstance(pos, dict)]
    display_positions = sorted(
        valid_positions,
        key=lambda pos: abs(_position_market_value(pos) or Decimal("0")),
        reverse=True,
    )

    lines.append("")
    lines.append(f"🧾 <b>Позиции ({len(valid_positions)})</b>")
    for i, pos in enumerate(display_positions[:12], start=1):
        ticker = str(pos.get("ticker") or pos.get("figi") or "UNKNOWN")
        instrument_type = _portfolio_type_label(pos.get("instrumentType"))
        quantity_lots = _format_decimal_plain(_money_from_quotation(pos.get("quantityLots")))
        quantity_total = _format_decimal_plain(_money_from_quotation(pos.get("quantity")))
        pos_currency = _extract_currency(pos.get("currentPrice"), currency)
        avg_price = _format_money_human(pos.get("averagePositionPrice"), fallback_currency=pos_currency)
        current_price = _format_money_human(pos.get("currentPrice"), fallback_currency=pos_currency)
        pos_yield_raw = pos.get("expectedYield")
        pos_yield_dec = _decimal_from_any(pos_yield_raw)
        pos_yield_percent = _position_yield_percent(pos)
        market_value = _position_market_value(pos)
        weight_text = "—"
        if market_value is not None and total_all > 0:
            weight_text = f"{_format_decimal_human((market_value / total_all * Decimal('100')).quantize(Decimal('0.01')), decimals=2)}%"
        market_value_text = f"{_format_decimal_human(market_value)} {pos_currency}" if market_value is not None else "—"
        blocked = bool(pos.get("blocked"))
        blocked_lots = _format_decimal_plain(_money_from_quotation(pos.get("blockedLots")))
        lines.append(f"{i}. <b>{html.escape(ticker)}</b> · {html.escape(instrument_type)}")
        lines.append(f"   Стоимость: <b>{market_value_text}</b> · Доля: <b>{weight_text}</b>")
        lines.append(f"   Количество: <b>{quantity_total}</b> · Лоты: <b>{quantity_lots}</b>")
        lines.append(f"   Цена: <b>{avg_price}</b> → <b>{current_price}</b>")
        lines.append(f"   {_pl_badge(pos_yield_dec)} P/L: <b>{_format_signed_money(pos_yield_dec, pos_currency)}</b> · {_format_signed_percent(pos_yield_percent)}")
        lines.append(f"   Блокировка: <b>{'есть' if blocked else 'нет'}</b> · лотов: <b>{blocked_lots}</b>")
    if len(valid_positions) > 12:
        lines.append("")
        lines.append(f"Показаны первые <b>12</b> из <b>{len(valid_positions)}</b> позиций по размеру позиции.")
    return "\n".join(lines)


def _render_operations(details: dict, days: int) -> str:
    if not isinstance(details, dict):
        return "🕘 <b>Операции</b>\nДанные по операциям пока недоступны."
    operations = details.get("operations")
    if not isinstance(operations, list):
        return "🕘 <b>Операции</b>\nДанные по операциям пока недоступны."

    total_in = Decimal("0")
    total_out = Decimal("0")
    for op in operations:
        if not isinstance(op, dict):
            continue
        payment_dec = _decimal_from_any(op.get("payment"))
        if payment_dec is None:
            continue
        if payment_dec >= 0:
            total_in += payment_dec
        else:
            total_out += payment_dec

    lines = [
        f"🕘 <b>Операции за {days} дн.</b>",
        f"Всего: <b>{len(operations)}</b> · Вход: <b>{_format_decimal_human(total_in)} RUB</b> · Выход: <b>{_format_decimal_human(total_out)} RUB</b>",
        "",
    ]
    if not operations:
        lines.append("За выбранный период операций не найдено.")
        return "\n".join(lines)

    for i, op in enumerate(operations[:20], start=1):
        if not isinstance(op, dict):
            continue
        dt = str(op.get("date") or "—")
        try:
            parsed = datetime.fromisoformat(dt.replace("Z", "+00:00"))
            dt = parsed.strftime("%d.%m.%Y %H:%M:%S")
        except Exception:
            pass
        op_type = str(op.get("type") or op.get("operationType") or "UNKNOWN")
        op_type_human = op_type.replace("_", " ").strip() or "UNKNOWN"
        ticker = str(op.get("ticker") or op.get("figi") or "—")
        payment_raw = op.get("payment")
        payment_dec = _decimal_from_any(payment_raw)
        if isinstance(payment_raw, dict):
            payment = _format_money(payment_raw, fallback_currency="RUB")
        elif payment_dec is not None:
            payment = f"{_format_decimal_human(payment_dec)} RUB"
        else:
            payment = "—"
        flow_badge = "🟢" if (payment_dec or Decimal("0")) > 0 else "🔴" if (payment_dec or Decimal("0")) < 0 else "⚪"
        quantity = (
            _decimal_from_any(op.get("quantity"))
            or _decimal_from_any(op.get("quantityExecuted"))
            or _decimal_from_any(op.get("quantityLots"))
        )
        quantity_text = _format_decimal_human(quantity, decimals=6) if quantity is not None else "—"
        lines.append(f"{i}. <b>{html.escape(op_type_human)}</b> · <code>{html.escape(ticker)}</code>")
        lines.append(f"   {flow_badge} Сумма: <b>{html.escape(payment)}</b>")
        lines.append(f"   Количество: <b>{html.escape(quantity_text)}</b>")
        lines.append(f"   Дата: <b>{html.escape(dt)}</b>")
        lines.append("")
    if len(operations) > 20:
        lines.append(f"Показаны первые <b>20</b> из <b>{len(operations)}</b> операций.")
    return "\n".join(lines)


def _extract_active_orders_lists(orders_details: dict, stop_orders_details: dict) -> tuple[list[dict], list[dict]]:
    orders = orders_details.get("orders", []) if isinstance(orders_details, dict) else []
    stops = stop_orders_details.get("stopOrders", []) if isinstance(stop_orders_details, dict) else []
    safe_orders = [item for item in orders if isinstance(item, dict)] if isinstance(orders, list) else []
    safe_stops = [item for item in stops if isinstance(item, dict)] if isinstance(stops, list) else []
    return safe_orders, safe_stops


def _active_orders_keyboard(orders: list[dict], stops: list[dict]) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for i, order in enumerate(orders[:10], start=1):
        figi = str(order.get("figi") or order.get("instrumentId") or order.get("instrumentUid") or "—")
        rows.append(
            [
                InlineKeyboardButton(
                    text=f"❌ Отменить биржевую #{i} · {figi}",
                    callback_data=f"{AO_CB_CANCEL_ORDER_PREFIX}:{i}",
                )
            ]
        )

    for i, stop in enumerate(stops[:10], start=1):
        figi = str(stop.get("figi") or stop.get("instrumentId") or stop.get("instrumentUid") or "—")
        rows.append(
            [
                InlineKeyboardButton(
                    text=f"🛑 Отменить стоп #{i} · {figi}",
                    callback_data=f"{AO_CB_CANCEL_STOP_PREFIX}:{i}",
                )
            ]
        )

    rows.append([InlineKeyboardButton(text="🔄 Обновить", callback_data=AO_CB_REFRESH)])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _render_active_orders(orders_details: dict, stop_orders_details: dict) -> str:
    orders, stops = _extract_active_orders_lists(orders_details, stop_orders_details)

    lines = [
        "📂 <b>Активные заявки</b>",
        f"Биржевые: <b>{len(orders)}</b> · Стоп-приказы: <b>{len(stops)}</b>",
        "Для отмены нажмите кнопку с тем же номером заявки ниже.",
        "",
    ]
    if not orders:
        lines.append("Биржевых заявок сейчас нет.")
    else:
        lines.append("📈 <b>Биржевые заявки</b>")
        for i, order in enumerate(orders[:10], start=1):
            order_id = str(order.get("orderId") or order.get("order_id") or "—")
            figi = str(order.get("figi") or order.get("instrumentId") or order.get("instrumentUid") or "—")
            lines.append(f"{i}. FIGI: <code>{html.escape(figi)}</code>")
            lines.append(f"   ID: <code>{html.escape(order_id)}</code>")

    lines.append("")
    if not stops:
        lines.append("Стоп-приказов сейчас нет.")
    else:
        lines.append("🛑 <b>Стоп-приказы</b>")
        for i, stop in enumerate(stops[:10], start=1):
            stop_id = str(stop.get("stopOrderId") or stop.get("stop_order_id") or "—")
            figi = str(stop.get("figi") or stop.get("instrumentId") or stop.get("instrumentUid") or "—")
            lines.append(f"{i}. FIGI: <code>{html.escape(figi)}</code>")
            lines.append(f"   ID: <code>{html.escape(stop_id)}</code>")

    if len(orders) > 10 or len(stops) > 10:
        lines.append("")
        lines.append("Показаны первые 10 заявок каждого типа.")

    lines.append("")
    lines.append("Управление кнопками:")
    lines.append("• отмена биржевой заявки")
    lines.append("• отмена стоп-приказа")
    lines.append("• обновление списка")
    return "\n".join(lines)


def _monitor_list_keyboard(items: list[dict]) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for item in items:
        monitor_id = _parse_monitor_id(item.get("id"))
        if not monitor_id:
            continue
        ticker = str(item.get("ticker") or item.get("figi") or "UNKNOWN")
        status = "🟢" if item.get("is_active") else "⚪"
        interval = _format_interval_seconds(_extract_monitor_interval_seconds(item))
        rows.append(
            [
                InlineKeyboardButton(
                    text=f"{status} #{monitor_id} {ticker} · {interval}",
                    callback_data=f"{MON_CB_ITEM_PREFIX}:{monitor_id}",
                )
            ]
        )
    rows.append([InlineKeyboardButton(text="Обновить", callback_data=MON_CB_REFRESH)])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _monitor_detail_text(item: dict) -> str:
    status = "Включен" if item.get("is_active") else "Выключен"
    interval = _format_interval_seconds(_extract_monitor_interval_seconds(item))
    return (
        f"🧭 <b>Монитор #{_parse_monitor_id(item.get('id')) or '—'}</b>\n"
        f"Инструмент: <b>{html.escape(str(item.get('ticker') or item.get('figi') or '—'))}</b>\n"
        f"FIGI: <code>{html.escape(str(item.get('figi') or '—'))}</code>\n"
        f"Статус: <b>{status}</b>\n"
        f"Интервал: <b>{interval}</b>\n"
        f"Порог: <b>{_to_clean_num_str(item.get('threshold_percent'), '%')}</b> или <b>{_to_clean_num_str(item.get('threshold_rub'))} RUB</b>\n"
        f"Базовая цена: <b>{_to_clean_num_str(item.get('base_price'))} RUB</b>\n"
        f"Последняя проверка: <b>{_format_datetime(str(item.get('last_checked_at_msk')) if item.get('last_checked_at_msk') else None)}</b>\n"
        f"Последний алерт: <b>{_format_datetime(str(item.get('last_notified_at_msk')) if item.get('last_notified_at_msk') else None)}</b>\n\n"
        "Кнопки ниже:\n"
        "• Включить/выключить монитор\n"
        "• Изменить пороги (%) и RUB\n"
        "• Обновить базовую цену\n"
        "• Удалить монитор"
    )


def _monitor_detail_keyboard(item: dict) -> InlineKeyboardMarkup:
    monitor_id = _parse_monitor_id(item.get("id"))
    active = bool(item.get("is_active"))
    toggle_cb = f"{MON_CB_OFF_PREFIX}:{monitor_id}" if active else f"{MON_CB_ON_PREFIX}:{monitor_id}"
    toggle_text = "Выключить" if active else "Включить"
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text=toggle_text, callback_data=toggle_cb)],
            [InlineKeyboardButton(text="Изменить пороги", callback_data=f"{MON_CB_EDIT_PREFIX}:{monitor_id}")],
            [InlineKeyboardButton(text="Обновить базу", callback_data=f"{MON_CB_REBASE_PREFIX}:{monitor_id}")],
            [InlineKeyboardButton(text="Удалить", callback_data=f"{MON_CB_DEL_PREFIX}:{monitor_id}")],
            [InlineKeyboardButton(text="К списку", callback_data=MON_CB_LIST)],
        ]
    )


def _monitor_edit_menu_keyboard(monitor_id: int) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="Изменить %", callback_data=f"{MON_CB_EDIT_PERCENT_PREFIX}:{monitor_id}")],
            [InlineKeyboardButton(text="Изменить RUB", callback_data=f"{MON_CB_EDIT_RUB_PREFIX}:{monitor_id}")],
            [InlineKeyboardButton(text="Назад", callback_data=f"{MON_CB_EDIT_BACK_PREFIX}:{monitor_id}")],
        ]
    )


def _monitor_edit_menu_text(item: dict) -> str:
    monitor_id = _parse_monitor_id(item.get("id")) or "—"
    return (
        f"✏️ <b>Редактирование мониторинга #{monitor_id}</b>\n"
        f"Инструмент: <b>{html.escape(str(item.get('ticker') or item.get('figi') or '—'))}</b>\n"
        f"Текущий порог: <b>{_to_clean_num_str(item.get('threshold_percent'), '%')}</b> или <b>{_to_clean_num_str(item.get('threshold_rub'))} RUB</b>\n\n"
        "Выберите, что изменить:\n"
        "• <b>Изменить %</b> — порог относительного изменения\n"
        "• <b>Изменить RUB</b> — порог абсолютного изменения в рублях"
    )


def _profile_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="Обновить токен/account_id", callback_data=PROF_CB_UPDATE)],
            [InlineKeyboardButton(text="Проверить доступ", callback_data=PROF_CB_CHECK)],
        ]
    )
